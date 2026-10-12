from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.dml.color import RGBColor

ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = ROOT / 'docs/演示/交大ppt模板.pptx'
OUTPUT = ROOT / 'docs/演示/信息系统渗透智能化测试平台_甲方项目介绍_12页.pptx'
FONT = 'Microsoft YaHei'
NAVY = RGBColor(20, 48, 82)
BLUE = RGBColor(46, 111, 171)
TEXT = RGBColor(43, 54, 70)
MUTED = RGBColor(94, 108, 125)
GRID = RGBColor(210, 221, 234)
BG = RGBColor(250, 252, 255)
LIGHT = RGBColor(239, 247, 253)
WHITE = RGBColor(255, 255, 255)

def fill(s, c, t=0):
    s.fill.solid(); s.fill.fore_color.rgb = c; s.fill.transparency = t

def line_style(s, c=GRID, w=1, dash=None):
    s.line.color.rgb = c; s.line.width = Pt(w)
    if dash is not None: s.line.dash_style = dash

def rect(slide, x, y, w, h, fc=WHITE, lc=GRID, lw=.8, dash=None):
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    fill(s, fc); line_style(s, lc, lw, dash); return s

def txt(slide, value, x, y, w, h, size=18, color=TEXT, bold=False, align=PP_ALIGN.LEFT, font=FONT):
    s = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = s.text_frame; tf.clear(); tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(.05); tf.margin_top = tf.margin_bottom = Inches(.03)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]; p.alignment = align
    r = p.add_run(); r.text = value; r.font.name = font; r.font.size = Pt(size); r.font.bold = bold; r.font.color.rgb = color
    return s

def connector(slide, x1, y1, x2, y2, color=GRID, width=1, dash=None):
    s = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    line_style(s, color, width, dash); return s

def circle(slide, x, y, d, fc=WHITE, lc=BLUE, lw=1):
    s = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d)); fill(s, fc); line_style(s, lc, lw); return s

def arrow(slide, x, y, w=.26, h=.18, color=BLUE):
    s=slide.shapes.add_shape(MSO_SHAPE.CHEVRON, Inches(x), Inches(y), Inches(w), Inches(h))
    fill(s,color); line_style(s,color,.1); return s

def remove_ph(slide, idx):
    for sh in list(slide.placeholders):
        if sh.placeholder_format.idx == idx:
            e = sh._element; e.getparent().remove(e); return

def footer(slide, n):
    connector(slide, .55, 6.86, 12.78, 6.86, GRID, .7)
    txt(slide, '信息系统渗透智能化测试平台', .58, 6.94, 4.2, .18, 9, MUTED)
    txt(slide, f'{n:02d} / 12', 11.65, 6.94, 1.1, .18, 9, BLUE, True, PP_ALIGN.RIGHT, 'Arial')

def title(slide, n, kicker, heading, sub):
    txt(slide, kicker.upper(), .58, .22, 3.4, .2, 10, BLUE, True, font='Arial')
    txt(slide, heading, .58, .48, 11.7, .55, 30, NAVY, True)
    connector(slide, .58, 1.17, 12.76, 1.17, BLUE, 1.2)
    txt(slide, sub, .60, 1.28, 11.6, .3, 16, MUTED)
    footer(slide, n)

def bullet(slide, value, x, y, w, size=19):
    circle(slide, x, y+.12, .11, BLUE, BLUE, .2)
    txt(slide, value, x+.25, y, w-.25, .38, size, TEXT)

def placeholder(slide, x, y, w, h, number, label):
    s = rect(slide, x, y, w, h, BG, BLUE, 1.1, MSO_LINE_DASH_STYLE.DASH); s.name = f'图片占位 - 图{number} {label}'
    connector(slide, x+.25, y+.25, x+w-.25, y+h-.25, GRID, .6, MSO_LINE_DASH_STYLE.DASH)
    connector(slide, x+w-.25, y+.25, x+.25, y+h-.25, GRID, .6, MSO_LINE_DASH_STYLE.DASH)
    txt(slide, f'图 {number}', x+.2, y+h*.43, w-.4, .28, 19, BLUE, True, PP_ALIGN.CENTER)
    txt(slide, label, x+.2, y+h*.56, w-.4, .35, 14, MUTED, False, PP_ALIGN.CENTER)
    txt(slide, '后续替换为正式图片', x+.2, y+h-.35, w-.4, .18, 10, MUTED, False, PP_ALIGN.CENTER)

def blank(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6]); s.background.fill.solid(); s.background.fill.fore_color.rgb = BG; return s

def cover(prs):
    s = prs.slides.add_slide(prs.slide_layouts[11]); ph = {p.placeholder_format.idx:p for p in s.placeholders}
    tf = ph[13].text_frame; tf.clear(); tf.paragraphs[0].add_run().text = '项目介绍'
    for r in tf.paragraphs[0].runs: r.font.name='Arial'; r.font.size=Pt(16); r.font.bold=True; r.font.color.rgb=BLUE
    tf = ph[0].text_frame; tf.clear(); p=tf.paragraphs[0]; r=p.add_run(); r.text='信息系统渗透\n智能化测试平台'; r.font.name=FONT; r.font.size=Pt(32); r.font.bold=True; r.font.color.rgb=NAVY
    tf = ph[16].text_frame; tf.clear(); p=tf.paragraphs[0]; r=p.add_run(); r.text='面向信息系统安全测试的智能化平台\n建设目标 · 软硬件架构 · 核心能力 · 交付价值'; r.font.name=FONT; r.font.size=Pt(16); r.font.color.rgb=MUTED
    txt(s, '甲方项目介绍', 5.64, 5.42, 3.2, .28, 14, BLUE, True)
    txt(s, '项目方案与能力概览', 5.64, 5.82, 3.4, .22, 12, MUTED)

def build():
    prs = Presentation(str(TEMPLATE))
    ids = prs.slides._sldIdLst
    for sid in list(ids): prs.part.drop_rel(sid.rId); ids.remove(sid)
    cover(prs)

    s=blank(prs); title(s,2,'01 / 定位','项目建设目标与总体方案','以统一平台组织测试任务、工具资源与结果交付，形成安全测试业务闭环。')
    rect(s,.62,1.86,6.1,4.55,WHITE,GRID,.8); txt(s,'建设目标',.94,2.12,2,.3,22,NAVY,True)
    for i,v in enumerate(['统一管理测试目标、测试任务和测试资源','提升安全测试的自动化与智能化水平','形成扫描发现到验证报告的完整闭环','适配国产化软硬件与实验室、现场应用','实现过程留痕、证据关联和结果追溯']): bullet(s,v,.98,2.68+i*.60,5.4,18)
    placeholder(s,7.15,1.86,5.05,4.55,1,'平台总体方案或应用场景图')

    s=blank(prs); title(s,3,'02 / 架构','软硬件一体化架构','中心服务器集中承载平台能力，国产化云笔电作为多角色操作终端。')
    rect(s,.62,1.88,4.05,4.5,WHITE,BLUE,.9); txt(s,'硬件组成',.94,2.16,1.8,.3,22,NAVY,True)
    for i,(a,b) in enumerate([('中心服务器','海光 x86'),('云笔电客户端','飞腾 E2000Q ARM64'),('操作系统','银河麒麟 V10 SP1'),('扩展设备','GPU / USB 硬件加密锁')]):
        y=2.80+i*.78; circle(s,.98,y+.06,.25,BLUE,BLUE,.2); txt(s,a,1.38,y,2.8,.22,16,TEXT,True); txt(s,b,1.38,y+.29,2.8,.18,12,MUTED)
    rect(s,4.98,1.88,4.15,4.5,WHITE,BLUE,.9); txt(s,'软件分层',5.30,2.16,1.8,.3,22,NAVY,True)
    for i,(a,b) in enumerate([('交互层','Qt5 / C++ 客户端'),('服务层','Node.js / 任务调度 / 权限'),('能力层','扫描 / 分析 / 预案 / 执行'),('资源层','SQLite / 容器 / 知识库 / 模型')]):
        y=2.80+i*.75; rect(s,5.32,y,1.12,.34,LIGHT,BLUE,.5); txt(s,a,5.35,y+.06,1.06,.16,10,BLUE,True,PP_ALIGN.CENTER); txt(s,b,6.72,y+.05,2.15,.2,12,TEXT)
    placeholder(s,9.45,1.88,2.75,4.5,2,'服务器、云笔电或部署现场照片')

    s=blank(prs); title(s,4,'03 / 部署','部署拓扑与运行模式','采用“中心服务 + 多端接入 + 测试环境”的集中部署方式。')
    rect(s,.62,1.88,7.6,4.5,WHITE,GRID,.8); rect(s,2.76,2.22,3.25,.7,LIGHT,BLUE,1); txt(s,'中心服务器｜海光 x86',2.88,2.36,3.0,.2,16,BLUE,True,PP_ALIGN.CENTER); txt(s,'后端 · 数据库 · 容器 · 模型服务',2.88,2.62,3,.15,10,MUTED,False,PP_ALIGN.CENTER); connector(s,4.39,2.92,4.39,3.38,BLUE,1.2); txt(s,'PoE 局域网',3.72,3.05,1.35,.18,12,BLUE,True,PP_ALIGN.CENTER); connector(s,1.25,3.48,7.55,3.48,BLUE,1)
    for x,a,b,c in [(0.88,'云笔电 1','操作端',BLUE),(3.15,'云笔电 2','控制端',BLUE),(5.42,'云笔电 3~8','扩展客户端',BLUE)]:
        connector(s,x+.76,3.48,x+.76,3.88,BLUE,.9); rect(s,x,3.88,1.52,.7,WHITE,BLUE,.8); txt(s,a,x+.05,4.02,1.42,.18,12,TEXT,True,PP_ALIGN.CENTER); txt(s,b,x+.05,4.28,1.42,.14,10,BLUE,False,PP_ALIGN.CENTER)
    connector(s,4.39,4.58,4.39,4.98,BLUE,1); rect(s,3.08,4.98,2.62,.62,LIGHT,BLUE,.8); txt(s,'测试靶场 / 被测系统',3.20,5.18,2.38,.16,13,BLUE,True,PP_ALIGN.CENTER); txt(s,'任务与数据统一经中心服务器转发。',1.05,5.86,6.75,.2,12,MUTED,False,PP_ALIGN.CENTER)
    placeholder(s,8.58,1.88,3.62,4.5,3,'网络部署拓扑图或现场连接照片')

    s=blank(prs); title(s,5,'04 / 流程','核心业务流程','围绕正式测试任务，形成从授权范围确认到测试成果交付的闭环管理。')
    steps=[('01','任务准备','目标 / 授权'),('02','扫描发现','资产 / 服务'),('03','智能分析','证据 / 知识'),('04','预案制定','生成 / 审核'),('05','执行验证','受控执行'),('06','结果评估','报告 / 复测')]
    for i,(n,a,b) in enumerate(steps):
        row=0 if i<3 else 1; col=i if i<3 else i-3; x=.78+col*2.58; y=2.0+row*1.54; rect(s,x,y,2.14,.96,WHITE,BLUE,.9); circle(s,x+.16,y+.18,.34,BLUE,BLUE,.2); txt(s,n,x+.16,y+.29,.34,.1,8,WHITE,True,PP_ALIGN.CENTER,'Arial'); txt(s,a,x+.65,y+.18,1.32,.2,14,TEXT,True); txt(s,b,x+.18,y+.61,1.78,.15,10,MUTED,False,PP_ALIGN.CENTER);
        if col<2: arrow(s,x+2.23,y+.40,.22,.16,BLUE)
    connector(s,3.96,2.98,3.96,3.33,BLUE,1); arrow(s,3.84,3.18,.24,.16,BLUE); rect(s,.78,5.28,7.18,.72,LIGHT,BLUE,.6); txt(s,'任务上下文贯穿全过程：补充扫描、重新分析、预案修订和结果复核持续关联。',1.0,5.51,6.75,.2,13,BLUE,True,PP_ALIGN.CENTER); placeholder(s,8.48,1.96,3.72,4.04,4,'任务流程或任务详情界面截图')

    s=blank(prs); title(s,6,'05 / 能力','平台功能模块与资源体系','七个一级模块覆盖测试组织、专业分析、资源沉淀、成果管理和运行保障。')
    mods=[('工作台','全局态势、待办与运行状态'),('测试任务','目标配置与流程管理'),('资产与拓扑','资产、服务与网络关系'),('智能分析','风险识别与证据关联'),('预案与知识库','Playbook、知识图谱、工具'),('结果与报告','验证结果与报告导出'),('系统管理','用户、权限、模型与审计')]
    for i,(a,b) in enumerate(mods):
        x=.72+(i%2)*3.48; y=1.98+(i//2)*.91; rect(s,x,y,3.15,.64,WHITE,BLUE,.7); txt(s,a,x+.18,y+.15,1.28,.18,14,BLUE,True); txt(s,b,x+1.48,y+.12,1.48,.28,11,MUTED)
    rect(s,.72,5.62,6.55,.68,LIGHT,BLUE,.6); txt(s,'资源体系：网络探测 · 漏洞扫描 · Web 安全 · 凭证测试 · 漏洞验证 · 云安全',.94,5.84,6.1,.18,11,BLUE,True,PP_ALIGN.CENTER); placeholder(s,8.02,1.98,4.18,4.32,5,'工作台、资产拓扑或功能模块界面截图')

    s=blank(prs); title(s,7,'06 / 智能协作','AI 智能协作能力','AI 连接扫描事实、知识资源、结构化预案和测试报告，辅助人工完成分析与决策。')
    rect(s,.68,1.92,7.08,4.35,WHITE,BLUE,.9); txt(s,'协作链路',.98,2.20,1.6,.3,22,NAVY,True)
    for i,(x,a,b) in enumerate([(1.05,'扫描事实','资产、端口、服务、发现'),(3.28,'智能分析','风险候选、证据、建议'),(5.51,'预案与报告','方案生成、结果整理')]):
        circle(s,x,2.96,1.26,WHITE,BLUE,1.2); txt(s,str(i+1),x+.47,3.18,.32,.16,12,BLUE,True,PP_ALIGN.CENTER,'Arial'); txt(s,a,x+.08,3.55,1.10,.18,12,TEXT,True,PP_ALIGN.CENTER); txt(s,b,x-.25,4.37,1.76,.25,10,MUTED,False,PP_ALIGN.CENTER);
        if i<2: arrow(s,x+1.45,3.50,.38,.2,BLUE)
    rect(s,1.02,5.30,6.35,.72,LIGHT,BLUE,.6); txt(s,'AI 辅助决策  →  人工审核确认  →  系统受控执行',1.26,5.53,5.88,.2,15,BLUE,True,PP_ALIGN.CENTER); placeholder(s,8.08,1.92,4.12,4.35,6,'AI 协作、预案编辑或智能分析界面截图')

    s=blank(prs); title(s,8,'07 / 安全','权限、安全与审计保障','以授权边界、人工审核、执行留痕和证据追溯保障测试过程安全可控。')
    rect(s,.68,1.92,7.08,4.35,WHITE,BLUE,.9); txt(s,'安全控制链',.98,2.20,1.8,.3,22,NAVY,True)
    for i,(a,b) in enumerate([('身份与权限','用户、角色和权限分级管理'),('授权与审批','测试范围和高风险操作控制'),('过程留痕','工具调用、日志和原始输出记录'),('证据追溯','任务、预案、执行和报告关联')]):
        y=2.82+i*.72; circle(s,1.04,y+.04,.25,BLUE,BLUE,.2); txt(s,'✓',1.04,y+.10,.25,.10,9,WHITE,True,PP_ALIGN.CENTER); txt(s,a,1.48,y,1.48,.2,14,TEXT,True); txt(s,b,3.12,y,4.0,.2,12,MUTED)
    rect(s,1.0,5.80,6.34,.34,LIGHT,BLUE,.5); txt(s,'可结合 USB 硬件加密锁实现软件授权控制。',1.20,5.88,5.94,.14,10,BLUE,True,PP_ALIGN.CENTER); placeholder(s,8.08,1.92,4.12,4.35,7,'权限管理、审批或审计日志界面截图')

    s=blank(prs); title(s,9,'08 / 交付','结果评估与报告交付','将测试过程转化为可审核、可追溯、可交付的结构化成果。')
    placeholder(s,.68,1.92,4.40,4.35,8,'测试报告封面、结果总览或报告页面截图'); rect(s,5.45,1.92,6.75,4.35,WHITE,BLUE,.9); txt(s,'甲方可获得的测试成果',5.78,2.20,3.0,.3,22,NAVY,True)
    for i,v in enumerate(['测试任务总体概览与资产、服务清单','风险发现及经过验证的测试结果','攻击步骤、工具执行记录与网络数据','原始输出、日志、证据和风险评分','结构化测试报告、审核、导出和归档文件']): bullet(s,v,5.84,2.84+i*.55,5.78,16)
    rect(s,5.78,5.82,5.98,.34,LIGHT,BLUE,.5); txt(s,'报告 → 任务 → 预案 → 执行记录 → 原始证据',5.98,5.90,5.58,.14,11,BLUE,True,PP_ALIGN.CENTER)

    s=blank(prs); title(s,10,'09 / 应用','典型应用场景与应用价值','面向实验室、教学实训、国产化适配和现场检测等多类安全测试应用。')
    for i,(a,b) in enumerate([('实验室测试','网络安全测试与环境验证'),('国产化适配','软硬件环境兼容性测试'),('云笔电协同','多角色任务协同与控制'),('靶场实训','教学、演练与能力评估'),('现场检测','授权范围内的安全检测'),('结果复盘','报告归档与安全改进')]):
        x=.72+(i%2)*3.42; y=1.98+(i//2)*.98; rect(s,x,y,3.08,.76,WHITE,BLUE,.7); txt(s,a,x+.22,y+.15,1.35,.2,14,BLUE,True); txt(s,b,x+.22,y+.44,2.55,.15,10,MUTED)
    rect(s,.72,5.38,6.60,.80,LIGHT,BLUE,.6); txt(s,'应用价值',1.02,5.64,1.0,.2,15,BLUE,True); txt(s,'提升效率 · 规范流程 · 降低门槛 · 强化可信度与可追溯性',2.10,5.62,4.82,.2,12,TEXT,True,PP_ALIGN.CENTER); placeholder(s,8.00,1.98,4.20,4.20,9,'实验室、云笔电协同或现场测试照片')

    s=blank(prs); title(s,11,'10 / 交付','部署运维、扩展与交付内容','支持集中部署与持续扩展，并形成软件、资源、文档和环境配置等交付物。')
    rect(s,.68,1.92,5.56,3.98,WHITE,BLUE,.9); txt(s,'部署与扩展能力',.98,2.20,2.5,.3,22,NAVY,True)
    for i,v in enumerate(['服务器集中部署，客户端批量接入','支持离线安装、物理迁移和镜像更新','适配国产操作系统与多种终端环境','可扩展工具、模型、靶场与专用硬件']): bullet(s,v,1.02,2.88+i*.60,4.8,17)
    rect(s,6.48,1.92,5.72,3.98,WHITE,BLUE,.9); txt(s,'项目交付内容',6.78,2.20,2.1,.3,22,NAVY,True)
    for i,v in enumerate(['平台客户端与中心服务端软件','数据库、数据管理模块与工具镜像','AI 模型接入、任务与报告管理能力','部署、用户、运维与硬件网络资料']): bullet(s,v,6.82,2.88+i*.60,4.9,17)
    placeholder(s,.68,6.10,11.52,.58,10,'硬件设备、软件安装包或交付清单示意图')

    s=blank(prs); txt(s,'信息系统渗透智能化测试平台',1.30,1.62,10.7,.34,23,BLUE,True,PP_ALIGN.CENTER); txt(s,'构建统一、智能、可追溯的安全测试工作平台',1.1,2.34,11.1,.65,34,NAVY,True,PP_ALIGN.CENTER); connector(s,4.26,3.28,9.07,3.28,BLUE,1.2); txt(s,'软硬件协同部署 · 全流程任务闭环 · 人机协作分析 · 结果证据追溯',1.1,3.68,11.1,.26,16,MUTED,False,PP_ALIGN.CENTER); txt(s,'谢谢各位领导、老师，欢迎指导',1.1,4.48,11.1,.36,22,TEXT,True,PP_ALIGN.CENTER); txt(s,'THANK YOU',1.1,5.18,11.1,.22,12,BLUE,True,PP_ALIGN.CENTER,'Arial'); footer(s,12)
    prs.save(str(OUTPUT)); print(OUTPUT)

if __name__ == '__main__': build()
