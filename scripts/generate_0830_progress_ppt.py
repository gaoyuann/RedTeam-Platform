from pathlib import Path
from pptx import Presentation
from pptx.util import Inches
from pptx.enum.text import PP_ALIGN

from generate_redteam_intro_ppt import (
    W, H, BG, PANEL, PANEL_2, TEXT, MUTED, CYAN, BLUE, PURPLE, GREEN, AMBER, RED, GRID,
    MONO, set_fill, set_line, add_rect, add_text, add_line, add_arrow, add_circle,
    add_chip, add_header, add_dark_bg, draw_shield,
)


def status(slide, label, x, y, color=GREEN, width=0.72):
    add_chip(slide, label, x, y, width, color, size=8)


def make_deck():
    prs = Presentation()
    prs.slide_width = Inches(W)
    prs.slide_height = Inches(H)
    blank = prs.slide_layouts[6]

    # 1. Cover
    slide = prs.slides.add_slide(blank); add_dark_bg(slide, grid=False)
    add_text(slide, "阶段工作汇报", 0.78, 1.18, 3.8, 0.56, size=30, color=TEXT, bold=True)
    add_text(slide, "2026.08.30 — 2026.09.03", 0.82, 1.88, 3.8, 0.24, size=12, color=CYAN, bold=True, font=MONO)
    add_text(slide, "信息系统渗透智能化测试平台", 0.82, 2.54, 5.80, 0.34, size=17, color=TEXT, bold=True)
    add_text(slide, "RedTeam-Platform", 0.82, 2.98, 4.8, 0.28, size=15, color=MUTED, font=MONO)
    add_text(slide, "本阶段围绕 8 月 30 日汇报提出的问题，集中完善攻击编排、扫描工作流、评估输出和演示页面。", 0.82, 3.56, 5.72, 0.60, size=13, color=MUTED)
    for i, (value, label, col) in enumerate([("4", "工作主线", CYAN), ("3", "攻击阶段", RED), ("4", "应用目标类型", PURPLE), ("2", "新增闭环引擎", GREEN)]):
        x = 0.82 + i * 1.50
        add_rect(slide, x, 5.16, 1.24, 0.78, fill=PANEL, line=GRID, lw=0.6)
        add_text(slide, value, x + 0.10, 5.28, 0.44, 0.30, size=20, color=col, bold=True, font=MONO)
        add_text(slide, label, x + 0.56, 5.33, 0.58, 0.28, size=8, color=MUTED)
    # right visual: four blocks converging to demo
    labels = [("CAMPAIGN", CYAN), ("PIPELINE", PURPLE), ("REPORT", GREEN), ("TOPOLOGY", AMBER)]
    for i, (lab, col) in enumerate(labels):
        x = 8.18 + (i % 2) * 1.90; y = 2.20 + (i // 2) * 1.26
        add_rect(slide, x, y, 1.58, 0.68, fill=PANEL, line=col, lw=0.8)
        add_text(slide, lab, x, y + 0.22, 1.58, 0.18, size=9, color=col, bold=True, align=PP_ALIGN.CENTER, font=MONO)
    add_line(slide, 9.94, 3.20, 10.66, 3.20, color=GRID, width=1.2)
    add_line(slide, 10.66, 3.20, 10.66, 4.16, color=GRID, width=1.2)
    add_line(slide, 10.66, 4.16, 10.20, 4.16, color=GRID, width=1.2)
    draw_shield(slide, 10.22, 3.72, 1.15, 1.30, CYAN)
    add_text(slide, "DEMO", 10.33, 4.15, 0.92, 0.18, size=9, color=TEXT, bold=True, align=PP_ALIGN.CENTER, font=MONO)
    add_text(slide, "WORK REPORT  ·  REDTEAM-PLATFORM", 7.72, 6.82, 4.62, 0.18, size=8, color=MUTED, align=PP_ALIGN.RIGHT, font=MONO)

    # 2. Scope and work packages
    slide = prs.slides.add_slide(blank); add_dark_bg(slide); add_header(slide, "01 / SCOPE", "8 月 30 日之后：围绕四类问题集中补功能", 2)
    add_text(slide, "本阶段以“功能能跑、链路串起来、结果看得懂”为验收导向", 0.62, 1.46, 8.0, 0.28, size=13, color=MUTED)
    work = [
        ("01", "攻击测试编排", "把三类攻击从标签筛选提升为可管理的阶段化战役", CYAN),
        ("02", "扫描到攻击串联", "补充应用扫描、目标分类、分析和一站式流水线能力", PURPLE),
        ("03", "评估与报告输出", "补齐报告预览/导出、网络捕获和分析入口", GREEN),
        ("04", "系统与演示体验", "补齐权限、靶场、拓扑、总览大屏和交互细节", AMBER),
    ]
    for i, (num, title, desc, col) in enumerate(work):
        x = 0.72 + (i % 2) * 6.00; y = 2.05 + (i // 2) * 1.62
        add_rect(slide, x, y, 5.54, 1.28, fill=PANEL, line=col, lw=0.8)
        add_circle(slide, x + 0.24, y + 0.32, 0.40, fill=col, line=col, lw=0.4)
        add_text(slide, num, x + 0.24, y + 0.45, 0.40, 0.14, size=9, color=BG, bold=True, align=PP_ALIGN.CENTER, font=MONO)
        add_text(slide, title, x + 0.88, y + 0.23, 2.20, 0.22, size=13, color=TEXT, bold=True)
        add_text(slide, desc, x + 0.88, y + 0.60, 4.16, 0.34, size=10, color=MUTED)
    add_rect(slide, 0.72, 5.60, 11.54, 0.82, fill=PANEL_2, line=GRID, lw=0.6)
    add_text(slide, "时间线", 1.02, 5.88, 0.72, 0.18, size=10, color=CYAN, bold=True, font=MONO)
    for i, (date, lab, col) in enumerate([("08.30", "问题拆解", CYAN), ("09.02", "扫描交互优化", PURPLE), ("09.03", "战役/大屏/拓扑联动", GREEN)]):
        x = 2.18 + i * 3.06
        add_circle(slide, x, 5.82, 0.22, fill=col, line=col, lw=0.3)
        add_text(slide, date, x + 0.36, 5.78, 0.82, 0.18, size=9, color=col, bold=True, font=MONO)
        add_text(slide, lab, x + 1.18, 5.78, 1.36, 0.18, size=9, color=TEXT, bold=True)
        if i < 2: add_arrow(slide, x + 2.36, 5.85, 0.36, 0.14, color=col, transparency=15)

    # 3. Campaign
    slide = prs.slides.add_slide(blank); add_dark_bg(slide); add_header(slide, "02 / CAMPAIGN", "攻击战役：三类攻击进入阶段化管理", 3)
    add_text(slide, "新增 Campaign 数据模型、后端接口和 Qt 页面，支持多 Playbook 编排", 0.62, 1.46, 9.0, 0.28, size=13, color=MUTED)
    phases = [("01", "数据抵近窃取", "凭据 / 文件 / 数据", CYAN), ("02", "信息篡改欺骗", "DNS / 会话 / 误导", PURPLE), ("03", "关键设备夺控", "横向 / 控制 / 持久", RED)]
    for i, (num, title, desc, col) in enumerate(phases):
        x = 0.84 + i * 3.04
        add_rect(slide, x, 2.10, 2.54, 1.30, fill=PANEL, line=col, lw=0.9)
        add_circle(slide, x + 0.20, 2.35, 0.36, fill=col, line=col, lw=0.4)
        add_text(slide, num, x + 0.20, 2.47, 0.36, 0.14, size=8, color=BG, bold=True, align=PP_ALIGN.CENTER, font=MONO)
        add_text(slide, title, x + 0.70, 2.33, 1.50, 0.24, size=12, color=TEXT, bold=True)
        add_text(slide, desc, x + 0.20, 2.92, 2.12, 0.20, size=9, color=MUTED)
        if i < 2: add_arrow(slide, x + 2.60, 2.67, 0.38, 0.18, color=col, transparency=10)
    # model / UI cards
    cards = [(0.84, "数据模型", "campaigns / phases / playbooks / artifacts", CYAN), (4.02, "前端页面", "战役列表 + 阶段指示 + 阶段产物", PURPLE), (7.20, "运行控制", "启动 / 暂停 / 终止 / 查看战役报告", GREEN), (10.38, "阶段执行", "串行 / 并行 / 条件模式", AMBER)]
    for x, title, desc, col in cards:
        add_rect(slide, x, 4.15, 2.54, 1.36, fill=PANEL, line=col, lw=0.7)
        add_text(slide, title, x + 0.18, 4.42, 1.70, 0.22, size=11, color=TEXT, bold=True)
        add_text(slide, desc, x + 0.18, 4.82, 2.12, 0.38, size=9, color=MUTED, font=MONO)
    add_rect(slide, 0.84, 5.92, 11.82, 0.48, fill=BG, line=GRID, lw=0.6)
    add_text(slide, "阶段产物可跨阶段传递：例如凭据、发现 IP、访问令牌进入后续 Playbook。", 1.08, 6.06, 11.30, 0.18, size=10, color=CYAN, bold=True, align=PP_ALIGN.CENTER)

    # 4. Pipeline
    slide = prs.slides.add_slide(blank); add_dark_bg(slide); add_header(slide, "03 / PIPELINE", "扫描流水线：补上分析层，形成可追踪的状态机", 4)
    add_text(slide, "新增 Pipeline 表、路由和执行引擎，状态与日志通过 WebSocket 实时广播", 0.62, 1.46, 9.6, 0.28, size=13, color=MUTED)
    steps = [("SCAN", "多类型扫描", CYAN), ("ANALYZE", "LLM 分析技术栈 / 攻击面", BLUE), ("GENERATE", "匹配或生成 Playbook", PURPLE), ("EXECUTE", "启动执行并记录证据", RED), ("REPORT", "汇总结果", GREEN)]
    for i, (lab, desc, col) in enumerate(steps):
        x = 0.78 + i * 2.48
        add_rect(slide, x, 2.35, 1.92, 1.26, fill=PANEL, line=col, lw=0.8)
        add_text(slide, lab, x, 2.62, 1.92, 0.20, size=10, color=col, bold=True, align=PP_ALIGN.CENTER, font=MONO)
        add_text(slide, desc, x + 0.16, 3.00, 1.60, 0.32, size=9, color=MUTED, align=PP_ALIGN.CENTER)
        if i < 4: add_arrow(slide, x + 1.98, 2.88, 0.34, 0.16, color=col, transparency=10)
    # state strip
    add_text(slide, "状态轨迹", 0.82, 4.36, 1.02, 0.20, size=10, color=CYAN, bold=True, font=MONO)
    states = [("created", GRID), ("scanning", CYAN), ("analyzing", BLUE), ("generating", PURPLE), ("ready", GREEN), ("executing", RED), ("completed", GREEN)]
    for i, (lab, col) in enumerate(states):
        x = 2.08 + i * 1.46
        add_circle(slide, x, 4.30, 0.24, fill=col, line=col, lw=0.3)
        add_text(slide, lab, x - 0.22, 4.68, 0.72, 0.16, size=7, color=TEXT, align=PP_ALIGN.CENTER, font=MONO)
        if i < len(states) - 1: add_arrow(slide, x + 0.34, 4.36, 0.30, 0.12, color=col, transparency=20)
    add_rect(slide, 0.82, 5.30, 11.72, 0.82, fill=PANEL, line=GRID, lw=0.6)
    add_text(slide, "已补充", 1.10, 5.58, 0.70, 0.18, size=10, color=GREEN, bold=True)
    add_text(slide, "取消 / 重试 / 失败记录  ·  分析结果持久化  ·  生成 Playbook 关联  ·  pipeline 日志推送", 2.02, 5.54, 9.86, 0.22, size=10, color=TEXT, font=MONO)

    # 5. Scan UX and application targets
    slide = prs.slides.add_slide(blank); add_dark_bg(slide); add_header(slide, "04 / SCAN UX", "扫描页面：从“点一下跳走”改为“先看清再执行”", 5)
    add_text(slide, "扫描结果、推荐预案和 AI 生成预案集中在同一页面完成确认", 0.62, 1.46, 8.6, 0.28, size=13, color=MUTED)
    add_rect(slide, 0.72, 2.06, 4.02, 3.98, fill=PANEL, line=CYAN, lw=0.8)
    add_text(slide, "应用目标识别", 1.02, 2.36, 1.56, 0.22, size=13, color=TEXT, bold=True)
    for i, (lab, desc, col) in enumerate([("Web 应用", "通用站点与服务", CYAN), ("REST API", "接口端点与认证", PURPLE), ("GraphQL API", "Schema 与查询面", AMBER), ("SPA", "前端路由与资源", GREEN)]):
        y = 2.92 + i * 0.58
        add_circle(slide, 1.04, y + 0.05, 0.18, fill=col, line=col, lw=0.3)
        add_text(slide, lab, 1.40, y, 1.25, 0.20, size=10, color=TEXT, bold=True)
        add_text(slide, desc, 2.82, y, 1.42, 0.20, size=9, color=MUTED)
    add_text(slide, "扫描后可动态修正目标类型", 1.02, 5.42, 2.80, 0.18, size=9, color=CYAN, bold=True, font=MONO)
    add_rect(slide, 5.08, 2.06, 7.58, 3.98, fill=PANEL, line=PURPLE, lw=0.8)
    add_text(slide, "交互链路", 5.38, 2.36, 1.20, 0.22, size=13, color=TEXT, bold=True)
    ux = [("推荐项", "点击仅选中", CYAN), ("详情", "查看步骤与说明", BLUE), ("AI 生成", "生成后先预览", PURPLE), ("前往执行", "确认后跳转", GREEN)]
    for i, (title, desc, col) in enumerate(ux):
        x = 5.40 + i * 1.76
        add_circle(slide, x, 3.22, 0.42, fill=col, line=col, lw=0.4)
        add_text(slide, str(i + 1), x, 3.36, 0.42, 0.14, size=9, color=BG, bold=True, align=PP_ALIGN.CENTER, font=MONO)
        add_text(slide, title, x - 0.16, 3.82, 0.76, 0.20, size=10, color=TEXT, bold=True, align=PP_ALIGN.CENTER)
        add_text(slide, desc, x - 0.44, 4.14, 1.30, 0.30, size=8, color=MUTED, align=PP_ALIGN.CENTER)
        if i < 3: add_arrow(slide, x + 0.62, 3.38, 0.44, 0.16, color=col, transparency=10)
    add_rect(slide, 5.40, 5.18, 6.70, 0.54, fill=PANEL_2, line=GRID, lw=0.6)
    add_text(slide, "whatweb/httpx · arjun/ffuf · nuclei 认证模板", 5.64, 5.34, 6.22, 0.18, size=9, color=PURPLE, bold=True, align=PP_ALIGN.CENTER, font=MONO)

    # 6. ReAct and evidence
    slide = prs.slides.add_slide(blank); add_dark_bg(slide); add_header(slide, "05 / REACT", "智能执行：把每一步结果变成下一步决策", 6)
    add_text(slide, "强化探索模式、启发式插入、失败转向和安全约束，并在前端展示推理证据", 0.62, 1.46, 10.2, 0.28, size=13, color=MUTED)
    add_rect(slide, 0.72, 2.04, 7.10, 4.06, fill=PANEL, line=RED, lw=0.8)
    loop = [("执行", CYAN), ("观察", BLUE), ("思考", PURPLE), ("行动", RED)]
    coords = [(1.14, 3.02), (3.22, 3.02), (3.22, 4.58), (1.14, 4.58)]
    for (lab, col), (x, y) in zip(loop, coords):
        add_rect(slide, x, y, 1.62, 0.62, fill=PANEL_2, line=col, lw=0.8)
        add_text(slide, lab, x, y + 0.19, 1.62, 0.20, size=12, color=TEXT, bold=True, align=PP_ALIGN.CENTER)
    add_arrow(slide, 2.82, 3.22, 0.34, 0.16, color=CYAN, transparency=10)
    add_arrow(slide, 3.88, 3.94, 0.16, 0.34, color=BLUE, transparency=10)
    add_arrow(slide, 2.82, 4.78, 0.34, 0.16, color=RED, transparency=10)
    add_arrow(slide, 1.78, 3.94, 0.16, 0.34, color=PURPLE, transparency=10)
    add_text(slide, "continue · adjust · insert · pivot · stop", 1.00, 5.58, 5.80, 0.18, size=10, color=RED, bold=True, font=MONO, align=PP_ALIGN.CENTER)
    add_rect(slide, 8.18, 2.04, 4.48, 4.06, fill=PANEL, line=PURPLE, lw=0.8)
    add_text(slide, "执行详情可视化", 8.48, 2.34, 1.90, 0.24, size=14, color=TEXT, bold=True)
    for i, (title, desc, col) in enumerate([("Cortex 决策面板", "Observation / Thought / Action", PURPLE), ("步骤执行表", "工具、参数、成功与输出摘要", CYAN), ("攻击证据", "载荷、证据记录与结果关联", GREEN), ("安全阀", "调用上限、黑名单、沙箱约束", AMBER)]):
        y = 2.98 + i * 0.62
        add_circle(slide, 8.52, y + 0.05, 0.16, fill=col, line=col, lw=0.3)
        add_text(slide, title, 8.84, y, 1.80, 0.20, size=10, color=TEXT, bold=True)
        add_text(slide, desc, 10.70, y, 1.54, 0.20, size=8, color=MUTED)

    # 7. Evaluation and capture
    slide = prs.slides.add_slide(blank); add_dark_bg(slide); add_header(slide, "06 / EVALUATION", "评估输出：报告与网络捕获入口同步补齐", 7)
    add_text(slide, "从执行记录、网络流量到最终报告，形成可回看的评估链路", 0.62, 1.46, 8.6, 0.28, size=13, color=MUTED)
    nodes = [("执行记录", "run / evidence", CYAN), ("网络捕获", "任务 / PCAP", PURPLE), ("捕获分析", "协议 / 会话 / 异常", BLUE), ("测试报告", "预览 / 导出", GREEN)]
    for i, (title, desc, col) in enumerate(nodes):
        x = 0.88 + i * 3.02
        add_rect(slide, x, 2.18, 2.36, 1.08, fill=PANEL, line=col, lw=0.8)
        add_text(slide, title, x + 0.18, 2.42, 1.72, 0.22, size=12, color=TEXT, bold=True)
        add_text(slide, desc, x + 0.18, 2.78, 1.92, 0.18, size=9, color=MUTED, font=MONO)
        if i < 3: add_arrow(slide, x + 2.44, 2.68, 0.42, 0.18, color=col, transparency=10)
    details = [(0.88, "报告预览", "JSON ↔ HTML 双模式", CYAN), (4.02, "多格式导出", "DOCX / PDF / HTML", GREEN), (7.16, "WPS 打开", "导出后直接打开编辑", PURPLE), (10.30, "捕获分析", "tcpdump / tshark 统计", AMBER)]
    for x, title, desc, col in details:
        add_rect(slide, x, 4.10, 2.36, 1.34, fill=PANEL, line=col, lw=0.7)
        add_text(slide, title, x + 0.18, 4.36, 1.80, 0.22, size=11, color=TEXT, bold=True)
        add_text(slide, desc, x + 0.18, 4.78, 1.92, 0.22, size=9, color=MUTED, font=MONO)
    add_rect(slide, 0.88, 5.82, 11.78, 0.54, fill=PANEL_2, line=GRID, lw=0.6)
    add_text(slide, "现状：前端入口、任务生命周期、分析服务和导出服务已补齐；目标环境仍需继续联调。", 1.12, 5.99, 11.30, 0.18, size=10, color=AMBER, bold=True, align=PP_ALIGN.CENTER)

    # 8. Topology + dashboard
    slide = prs.slides.add_slide(blank); add_dark_bg(slide); add_header(slide, "07 / VISUALIZATION", "拓扑与总览：把“发生了什么”展示出来", 8)
    add_text(slide, "新增拓扑探测/绘制交互与安全运营总览，便于演示时快速定位状态", 0.62, 1.46, 9.0, 0.28, size=13, color=MUTED)
    add_rect(slide, 0.72, 2.06, 5.68, 3.94, fill=PANEL, line=CYAN, lw=0.8)
    add_text(slide, "网络拓扑探测与绘制", 1.02, 2.34, 2.40, 0.24, size=14, color=TEXT, bold=True)
    add_text(slide, "扫描任务 → 智能生成拓扑 → 画布编辑 → 保存", 1.02, 2.76, 4.62, 0.22, size=10, color=CYAN, bold=True, font=MONO)
    # mini graph
    add_line(slide, 1.64, 4.48, 2.84, 3.62, color=GRID, width=1.0)
    add_line(slide, 2.84, 3.62, 4.18, 4.56, color=GRID, width=1.0)
    add_line(slide, 2.84, 3.62, 4.92, 3.40, color=GRID, width=1.0)
    for x, y, lab, col in [(1.40, 4.28, "GW", AMBER), (2.58, 3.36, "APP", CYAN), (3.94, 4.36, "DB", PURPLE), (4.68, 3.18, "API", GREEN)]:
        add_circle(slide, x, y, 0.48, fill=col, line=col, lw=0.4)
        add_text(slide, lab, x, y + 0.18, 0.48, 0.14, size=8, color=BG, bold=True, align=PP_ALIGN.CENTER, font=MONO)
    add_text(slide, "节点 / 连线 / 服务 / 状态均可编辑", 1.02, 5.34, 3.62, 0.18, size=9, color=MUTED)
    add_rect(slide, 6.72, 2.06, 5.94, 3.94, fill=PANEL, line=PURPLE, lw=0.8)
    add_text(slide, "安全运营总览", 7.02, 2.34, 1.60, 0.24, size=14, color=TEXT, bold=True)
    for i, (value, lab, col) in enumerate([("扫描", "任务统计", CYAN), ("执行", "实时活动", RED), ("报告", "结果沉淀", GREEN), ("连接", "WebSocket 状态", PURPLE)]):
        x = 7.02 + (i % 2) * 2.66; y = 2.98 + (i // 2) * 0.98
        add_rect(slide, x, y, 2.32, 0.72, fill=PANEL_2, line=col, lw=0.7)
        add_text(slide, value, x + 0.16, y + 0.16, 0.82, 0.22, size=12, color=col, bold=True)
        add_text(slide, lab, x + 1.02, y + 0.20, 1.10, 0.18, size=9, color=MUTED)
    add_text(slide, "总览页接入扫描 / 执行事件，页面可按角色显示。", 7.02, 5.40, 4.72, 0.18, size=9, color=PURPLE, bold=True)

    # 9. Management, labs and current status
    slide = prs.slides.add_slide(blank); add_dark_bg(slide); add_header(slide, "08 / DELIVERY", "管理与自检：让新增功能可配置、可验证", 9)
    add_text(slide, "本阶段同步补齐权限配置、靶场自检与页面联动，降低演示准备成本", 0.62, 1.46, 9.6, 0.28, size=13, color=MUTED)
    ops = [("权限矩阵", "可视化配置模块读写权限", CYAN), ("靶场管理", "启动 / 停止 / 自检工具链", GREEN), ("页面联动", "角色化导航与跨页刷新", PURPLE), ("安全约束", "输入校验、审计与权限守卫", AMBER)]
    for i, (title, desc, col) in enumerate(ops):
        x = 0.72 + (i % 2) * 6.00; y = 2.08 + (i // 2) * 1.40
        add_rect(slide, x, y, 5.54, 1.08, fill=PANEL, line=col, lw=0.8)
        add_circle(slide, x + 0.22, y + 0.29, 0.34, fill=col, line=col, lw=0.4)
        add_text(slide, str(i + 1), x + 0.22, y + 0.40, 0.34, 0.14, size=8, color=BG, bold=True, align=PP_ALIGN.CENTER, font=MONO)
        add_text(slide, title, x + 0.80, y + 0.20, 1.60, 0.22, size=12, color=TEXT, bold=True)
        add_text(slide, desc, x + 0.80, y + 0.56, 3.90, 0.20, size=10, color=MUTED)
    add_rect(slide, 0.72, 5.30, 11.54, 0.78, fill=PANEL_2, line=GRID, lw=0.6)
    add_text(slide, "演示准备链", 1.02, 5.58, 1.10, 0.18, size=10, color=CYAN, bold=True, font=MONO)
    chain = [("登录", CYAN), ("扫描", BLUE), ("推荐", PURPLE), ("执行", RED), ("报告", GREEN)]
    for i, (lab, col) in enumerate(chain):
        x = 2.36 + i * 1.78
        add_circle(slide, x, 5.52, 0.24, fill=col, line=col, lw=0.3)
        add_text(slide, lab, x + 0.38, 5.54, 0.82, 0.16, size=9, color=TEXT, bold=True)
        if i < 4: add_arrow(slide, x + 1.08, 5.58, 0.32, 0.12, color=col, transparency=15)
    add_text(slide, "当前状态：核心页面与后端能力已连接，正在持续做现场环境联调。", 0.72, 6.38, 11.54, 0.18, size=10, color=GREEN, bold=True, align=PP_ALIGN.CENTER)

    # 10. Summary
    slide = prs.slides.add_slide(blank); add_dark_bg(slide); add_header(slide, "09 / SUMMARY", "本阶段小结：功能从“分散可用”走向“连贯可演示”", 10)
    add_text(slide, "8 月 30 日以来，主要完成了三件事：补模型、串链路、做展示。", 0.62, 1.46, 9.0, 0.28, size=13, color=MUTED)
    summary = [("补模型", "Campaign / Pipeline / Capture 数据模型与 API", CYAN), ("串链路", "扫描 → 分析 → 生成 → 执行 → 报告", PURPLE), ("做展示", "ScanPage / Campaign / 拓扑 / 总览大屏", GREEN)]
    for i, (title, desc, col) in enumerate(summary):
        x = 0.84 + i * 4.08
        add_rect(slide, x, 2.22, 3.56, 1.56, fill=PANEL, line=col, lw=0.9)
        add_text(slide, title, x + 0.24, 2.52, 1.28, 0.24, size=14, color=TEXT, bold=True)
        add_text(slide, desc, x + 0.24, 2.98, 2.96, 0.44, size=10, color=MUTED)
    add_rect(slide, 0.84, 4.30, 11.82, 1.44, fill=PANEL, line=AMBER, lw=0.8)
    add_text(slide, "下一步重点", 1.14, 4.60, 1.40, 0.22, size=13, color=TEXT, bold=True)
    add_text(slide, "① 完成 Pipeline 一站式页面与端到端演示  ·  ② 继续验证抓包与评估链路  ·  ③ 固化拓扑和应用目标演示脚本", 2.72, 4.58, 9.18, 0.42, size=10, color=MUTED)
    add_text(slide, "汇报建议：按照“扫描 → 推荐 → 执行 → 证据 → 报告”主线现场演示。", 0.84, 6.24, 11.82, 0.22, size=12, color=CYAN, bold=True, align=PP_ALIGN.CENTER)

    return prs


if __name__ == "__main__":
    out_dir = Path(__file__).resolve().parents[1] / "docs" / "软件需规"
    out_dir.mkdir(parents=True, exist_ok=True)
    out_path = out_dir / "RedTeam-Platform_8.30以来工作汇报.pptx"
    make_deck().save(str(out_path))
    print(out_path)

