from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.dml.color import RGBColor

ROOT=Path(__file__).resolve().parents[1]
TEMPLATE=ROOT/'docs/演示/交大ppt模板.pptx'
OUTPUT=ROOT/'docs/演示/信息系统渗透智能化测试平台_甲方项目介绍_13页占位符版.pptx'
STANDARD_OUTPUT=ROOT/'docs/演示/信息系统渗透智能化测试平台_甲方项目介绍_13页.pptx'
FONT='Microsoft YaHei'; NAVY=RGBColor(15,20,35); TEXT=RGBColor(35,49,73); MUTED=RGBColor(96,111,132); BLUE=RGBColor(96,150,230); CYAN=RGBColor(88,182,229); GREEN=RGBColor(86,202,149); PURPLE=RGBColor(149,77,114); AMBER=RGBColor(255,186,85); GRID=RGBColor(222,230,240); BG=RGBColor(247,250,254); WHITE=RGBColor(255,255,255)
def fill(s,c,t=0): s.fill.solid(); s.fill.fore_color.rgb=c; s.fill.transparency=t
def border(s,c=GRID,w=.8,d=None): s.line.color.rgb=c; s.line.width=Pt(w); s.line.dash_style=d if d else s.line.dash_style
def rect(slide,x,y,w,h,fc=WHITE,lc=GRID,lw=.8,d=None):
    s=slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h)); fill(s,fc); border(s,lc,lw,d); return s
def txt(slide,v,x,y,w,h,size=16,c=TEXT,b=False,a=PP_ALIGN.LEFT,font=FONT):
    s=slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h)); tf=s.text_frame; tf.clear(); tf.word_wrap=True; tf.vertical_anchor=MSO_ANCHOR.MIDDLE; tf.margin_left=tf.margin_right=Inches(.05); tf.margin_top=tf.margin_bottom=Inches(.03); p=tf.paragraphs[0]; p.alignment=a; r=p.add_run(); r.text=v; r.font.name=font; r.font.size=Pt(size); r.font.bold=b; r.font.color.rgb=c; return s
def remove(slide,idx):
    for s in list(slide.placeholders):
        if s.placeholder_format.idx==idx: e=s._element; e.getparent().remove(e); return
def clear(prs):
    ids=prs.slides._sldIdLst
    for sid in list(ids): prs.part.drop_rel(sid.rId); ids.remove(sid)
def footer(slide,n): txt(slide,'信息系统渗透智能化测试平台',.58,7.04,4.3,.16,9,MUTED); txt(slide,f'{n:02d} / 13',11.55,7.04,1.25,.16,9,BLUE,True,PP_ALIGN.RIGHT,'Arial')
def placeholder(slide,x,y,w,h,code,purpose,color=BLUE):
    s=rect(slide,x,y,w,h,BG,color,1.1,MSO_LINE_DASH_STYLE.DASH); s.name=f'图片占位符 {code}（待采集）'; txt(slide,'图片占位符',x+.15,y+h*.30,w-.3,.25,17,color,True,PP_ALIGN.CENTER); txt(slide,f'{code}｜待采集',x+.15,y+h*.47,w-.3,.23,13,TEXT,True,PP_ALIGN.CENTER,'Arial'); txt(slide,purpose,x+.18,y+h*.62,w-.36,.40,12,MUTED,False,PP_ALIGN.CENTER); txt(slide,'正式截图替换后删除本行',x+.15,y+h-.28,w-.3,.14,9,MUTED,False,PP_ALIGN.CENTER)
def page(prs,n,title,kicker,conclusion):
    s=prs.slides.add_slide(prs.slide_layouts[14]);
    if s.shapes.title is not None:
        tf=s.shapes.title.text_frame; tf.clear(); r=tf.paragraphs[0].add_run(); r.text=title; r.font.name=FONT; r.font.size=Pt(30); r.font.bold=True; r.font.color.rgb=NAVY
    remove(s,10); txt(s,kicker.upper(),.42,.86,3.5,.16,10,BLUE,True,font='Arial'); txt(s,conclusion,.58,1.04,11.75,.68,20,MUTED); footer(s,n); return s
def panel(slide,x,y,w,h,title,items,color=BLUE):
    if h < 1:
        rect(slide,x,y,w,h)
        txt(slide,title,x+.15,y+.04,2.00,h-.08,18,NAVY,True)
        txt(slide,items[0],x+2.25,y+.04,w-2.45,h-.08,16,TEXT)
        return
    rect(slide,x,y,w,h)
    txt(slide,title,x+.24,y+.16,w-.48,.44,20,NAVY,True)
    compact = len(items) >= 6
    step=(h-.90)/len(items)
    for i,item in enumerate(items):
        if '|' in item:
            head, detail = item.split('|', 1)
        else:
            head, detail = item, ''
        base=y+.70+i*step
        if compact:
            txt(slide,head,x+.25,base,w-.50,.20,13.5,NAVY,True)
            if detail:
                txt(slide,detail,x+.25,base+.22,w-.50,max(.14,step-.24),10.5,MUTED)
            else:
                txt(slide,'',x+.25,base+.22,w-.50,max(.14,step-.24),10.5,MUTED)
        else:
            txt(slide,head,x+.25,base,w-.50,.22,16,NAVY,True)
            if detail:
                txt(slide,detail,x+.25,base+.27,w-.50,max(.18,step-.30),12.5,TEXT)
def cover(prs):
    s=prs.slides.add_slide(prs.slide_layouts[11]); p={x.placeholder_format.idx:x for x in s.placeholders}
    for idx,value,size,color,bold in [(13,'项目介绍',16,BLUE,True),(0,'智能测试平台与渗透测试平台',30,NAVY,True),(16,'软硬件一体化建设与功能演示',17,MUTED,False)]:
        tf=p[idx].text_frame; tf.clear(); r=tf.paragraphs[0].add_run(); r.text=value; r.font.name=FONT; r.font.size=Pt(size); r.font.color.rgb=color; r.font.bold=bold
    txt(s,'建设目标 · 软硬件架构 · 核心能力 · 交付价值',5.68,5.55,5.2,.26,13,MUTED); txt(s,'[文本占位-P01-C01：建设单位 / 汇报单位 / 日期]',5.68,5.96,5.2,.25,11,BLUE,True)
def directory(prs):
    s=prs.slides.add_slide(prs.slide_layouts[12]); remove(s,11)
    txt(s,'汇报目录',.68,.68,4.0,.50,30,NAVY,True)
    txt(s,'PROJECT OVERVIEW · SOFTWARE · CAPABILITIES · DELIVERY',.70,1.25,7.2,.20,10,BLUE,True,font='Arial')
    groups=[('01','项目与软件产品','第 3 页',BLUE),('02','软件架构与部署','第 4–5 页',CYAN),('03','平台模块与业务闭环','第 6 页',GREEN),('04','测试功能演示','第 7–11 页',PURPLE),('05','安全保障与建设价值','第 12–13 页',AMBER)]
    for i,(num,title,span,color) in enumerate(groups):
        y=1.78+i*.91
        rect(s,.82,y,11.55,.68,WHITE,GRID,.8)
        txt(s,num,.98,y+.14,.60,.22,15,color,True,PP_ALIGN.CENTER,'Arial')
        txt(s,title,1.88,y+.12,5.45,.24,18,NAVY,True)
        txt(s,span,9.90,y+.14,2.05,.20,12,MUTED,False,PP_ALIGN.RIGHT)
    txt(s,'从项目定位到软件架构，再到功能演示与交付价值',.82,6.55,11.55,.24,13,MUTED,False,PP_ALIGN.CENTER)
    footer(s,2)
def build():
    prs=Presentation(str(TEMPLATE)); clear(prs); cover(prs); directory(prs)
    s=page(prs,3,'项目概况与软件产品组成','01 / PRODUCT','两套软件共同支撑测试管理、智能分析和渗透验证。'); panel(s,.60,1.90,3.42,4.45,'项目建设对象',['信息系统与源代码软件系统|覆盖被测软件、服务和授权测试目标','目标靶场与授权测试范围|支持实验室、教学和现场场景','任务、资源、执行、证据、报告|把测试过程沉淀为可复核成果'],BLUE); panel(s,4.28,1.90,4.02,4.45,'智能测试平台',['任务组织与资源配置|统一安排目标、人员和测试资源','资产、拓扑与智能分析|从发现结果形成风险判断','知识、预案、评估与报告|支撑方案复用和成果交付'],CYAN); panel(s,8.56,1.90,4.16,4.45,'渗透测试平台',['漏洞扫描与目标验证|面向源代码系统和靶场执行检测','攻击载荷与受控执行|在授权范围内验证风险线索','证据采集与结果交付|保留原始输出并关联报告'],PURPLE)
    s=page(prs,4,'软件总体架构与技术分层','02 / SOFTWARE ARCHITECTURE','软件采用客户端、服务、能力引擎、数据与资源分层协同的架构。'); panel(s,.60,1.90,5.90,4.45,'四层架构',['交互层：Qt5/C++ 客户端|面向任务创建、监控、审阅和管理','服务层：Node.js、任务、权限、状态|统一编排接口、流程和访问控制','能力引擎：拓扑、扫描、分析、载荷、执行、报告|承载安全测试专业能力','数据资源：任务、证据、工具、知识、模型、靶场|提供可复用的测试资源'],BLUE); placeholder(s,6.82,1.90,5.90,4.45,'P04-A02','平台部署或服务状态页面',CYAN)
    s=page(prs,5,'软硬件组成与部署运行环境','03 / HARDWARE & DEPLOYMENT','中心服务、操作终端、测试网络和授权目标共同构成平台运行底座。'); placeholder(s,.60,1.90,7.30,4.45,'P05-H01','服务器、终端、靶场与加密锁实物组合图'); panel(s,8.20,1.90,4.52,4.45,'设备角色',['中心服务|承载平台、数据、工具和模型能力','操作终端|完成配置、过程监控和结果审阅','被测区|承载源代码系统与目标靶场','授权硬件|提供加密锁校验及现场扩展支撑'],CYAN)
    s=page(prs,6,'平台功能模块与业务闭环','04 / MODULES','平台模块围绕目标、任务、资源、执行、评估和报告组织完整业务。'); panel(s,.60,1.90,7.12,4.45,'七模块能力矩阵',['资源部署与系统配置|准备工具、靶场、模型和平台参数','预案与知识库|沉淀 Playbook、漏洞知识和工具方法','网络拓扑探测|识别主机、端口、服务及其关系','脆弱性扫描|形成扫描任务、风险线索和漏洞列表','漏洞攻击测试|选择目标、载荷和执行步骤进行验证','测试评估|汇总状态、评分、证据和报告','系统管理|维护用户、权限、任务和审计配置'],BLUE); panel(s,7.98,1.90,4.74,4.45,'业务闭环',['目标管理|明确对象和授权边界','任务组织|统一任务上下文和执行记录','资源调用|按目标匹配工具、模型和预案','受控执行|审核后按范围执行测试步骤','结果评估|关联证据并形成风险结论','报告交付|输出可复核、可归档的成果'],CYAN)
    s=page(prs,7,'被测源代码系统渗透与攻击测试','05 / SOURCE SYSTEMS','对每个被测源代码系统形成扫描发现、漏洞列表、逐项验证、定向攻击和证据交付闭环。'); placeholder(s,.60,1.90,7.12,3.72,'P07-S02','漏洞扫描结果列表'); panel(s,7.98,1.90,4.74,3.72,'演示要点',['按系统建账并确认范围|先核对目标、任务和授权边界','扫描结果形成漏洞列表|记录名称、等级、条件和状态','逐项验证并关联证据|保留验证输出和复核结论','攻击清单与目标保持关联|按候选漏洞逐项展示执行'],CYAN); panel(s,.60,5.92,12.12,.68,'证据字段',['目标标识 · 任务编号 · 漏洞名称 · 风险等级 · 验证状态 · 证据入口'],PURPLE)
    s=page(prs,8,'目标靶场渗透与攻击功能演示','06 / TARGET RANGE','对目标靶场通过任务统一组织目标配置、扫描、预案、受控执行、证据和报告。'); placeholder(s,.60,1.90,7.42,4.45,'P08-T04','任务执行工作台与关键输出'); panel(s,8.30,1.90,4.42,2.10,'目标与授权',['靶场名称与版本|说明演示对象和环境状态','目标地址与授权边界|限定测试范围和风险边界','任务 / Flow 编号|保证过程记录可回溯'],CYAN); placeholder(s,8.30,4.28,4.42,2.07,'P08-T05','报告或证据关联结果',PURPLE)
    s=page(prs,9,'PentAGI 与大语言模型载荷生成能力对照','07 / PAYLOAD GENERATION','两条路径均输入目标与漏洞线索，生成建议并记录审核、执行和证据状态。'); placeholder(s,.60,1.90,5.75,3.90,'P09-A01','PentAGI 生成页面'); placeholder(s,6.95,1.90,5.77,3.90,'P09-A03','大语言模型生成页面',CYAN); panel(s,.60,6.02,12.12,.68,'统一审核口径',['相同输入条件 · 来源标识 · 审核状态 · 允许执行状态 · 证据编号'],PURPLE)
    s=page(prs,10,'测试结果统计与证据追溯','08 / RESULTS','统计结果能够从统一口径下的汇总项回溯到任务、执行记录、漏洞和证据。'); placeholder(s,.60,1.90,7.35,4.45,'P10-R01','统计总览页面'); panel(s,8.25,1.90,4.47,4.45,'统计与追溯',['范围与时间窗口|明确数字对应的任务集合和时间段','任务状态与风险分布|说明总体状态和风险构成','验证状态与评分|展示结果质量和完成程度','统计项 → 任务 → 执行 → 证据 → 报告|支持逐级核验'],CYAN)
    s=page(prs,11,'网络拓扑探测与资产识别','09 / TOPOLOGY','拓扑探测将授权范围内的主机、端口和服务关系转化为可查看、可定位的资产视图。'); placeholder(s,.60,1.90,7.80,4.45,'P11-N03','拓扑探测结果主视图'); panel(s,8.70,1.90,4.02,4.45,'节点核验字段',['主机与地址|确认资产身份和授权范围','开放端口|定位可达服务入口','服务识别|支撑后续扫描和预案匹配','原始探测结果|保留复核依据','关联任务入口|将资产回到测试上下文'],CYAN)
    s=page(prs,12,'加密锁、安全控制与项目交付','10 / SECURITY & DELIVERY','加密锁、权限、过程审计和软件硬件交付资料共同支撑平台安全运行与项目交付。'); placeholder(s,.60,1.90,5.50,4.45,'P12-L01','加密锁实物与设备识别',AMBER); panel(s,6.42,1.90,6.30,4.45,'授权与交付',['授权状态与策略反馈|确认关键功能访问边界','用户、角色与审批控制|限定操作范围和高风险动作','日志、原始输出与证据关联|保证过程可审计、结果可复核','平台软件、硬件配置和配套文档|形成可持续运行的交付体系'],CYAN)
    s=prs.slides.add_slide(prs.slide_layouts[16])
    txt(s,'构建可验证、可管控、可追溯的安全测试能力',1.10,2.10,11.10,.70,29,NAVY,True,PP_ALIGN.CENTER)
    txt(s,'统一管理｜智能辅助｜受控验证｜证据追溯',1.10,3.05,11.10,.35,17,MUTED,False,PP_ALIGN.CENTER)
    txt(s,'[文本占位-P13-C01：正式致谢语或单位署名]',3.20,5.76,6.90,.25,12,BLUE,True,PP_ALIGN.CENTER); footer(s,13)
    for target in (OUTPUT, STANDARD_OUTPUT):
        prs.save(str(target))
    print(OUTPUT)
if __name__=='__main__': build()
