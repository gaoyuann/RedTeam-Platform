from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.dml.color import RGBColor

ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = ROOT / "docs" / "演示" / "交大ppt模板.pptx"
OUTPUT = ROOT / "docs" / "演示" / "信息系统渗透智能化测试平台_甲方项目介绍_12页.pptx"

FONT = "Microsoft YaHei"
FONT_EN = "Arial"
WHITE = RGBColor(255, 255, 255)
BG = RGBColor(249, 251, 254)
NAVY = RGBColor(15, 20, 35)
TEXT = RGBColor(35, 49, 73)
MUTED = RGBColor(96, 111, 132)
GRID = RGBColor(222, 230, 240)
BLUE = RGBColor(96, 150, 230)
CYAN = RGBColor(88, 182, 229)
PURPLE = RGBColor(149, 77, 114)
GREEN = RGBColor(86, 202, 149)
AMBER = RGBColor(255, 186, 85)
CORAL = RGBColor(241, 136, 112)
LIGHT_BLUE = RGBColor(241, 247, 253)
LIGHT_GREEN = RGBColor(241, 251, 247)
LIGHT_AMBER = RGBColor(255, 249, 239)
LIGHT_PURPLE = RGBColor(246, 243, 253)


def set_fill(shape, color, transparency=0):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.fill.transparency = transparency


def set_line(shape, color, width=1.0, transparency=0, dash=None):
    shape.line.color.rgb = color
    shape.line.width = Pt(width)
    shape.line.transparency = transparency
    if dash is not None:
        shape.line.dash_style = dash


def box(slide, x, y, w, h, fill=WHITE, line=GRID, lw=0.7, radius=True, transparency=0, dash=None):
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    set_fill(shape, fill, transparency)
    set_line(shape, line, lw, dash=dash)
    return shape


def add_text(slide, content, x, y, w, h, size=12, color=TEXT, bold=False,
             align=PP_ALIGN.LEFT, font=FONT, valign=MSO_ANCHOR.MIDDLE,
             margin=0.04, italic=False):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(margin)
    tf.margin_right = Inches(margin)
    tf.margin_top = Inches(margin)
    tf.margin_bottom = Inches(margin)
    tf.vertical_anchor = valign
    paragraph = tf.paragraphs[0]
    paragraph.alignment = align
    run = paragraph.add_run()
    run.text = content
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    return shape


def set_text(shape, content, size=18, color=TEXT, bold=False, align=PP_ALIGN.LEFT, font=FONT):
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    paragraph = tf.paragraphs[0]
    paragraph.alignment = align
    run = paragraph.add_run()
    run.text = content
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color


def add_line(slide, x1, y1, x2, y2, color=GRID, width=1.0, dash=None):
    line = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    set_line(line, color, width, dash=dash)
    return line


def add_circle(slide, x, y, d, fill=WHITE, line=BLUE, lw=1.0):
    shape = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d))
    set_fill(shape, fill)
    set_line(shape, line, lw)
    return shape


def add_chevron(slide, x, y, w=0.28, h=0.20, color=BLUE):
    shape = slide.shapes.add_shape(MSO_SHAPE.CHEVRON, Inches(x), Inches(y), Inches(w), Inches(h))
    set_fill(shape, color)
    set_line(shape, color, 0.1, transparency=100)
    return shape


def add_chip(slide, label, x, y, w, color=BLUE, size=8):
    shape = box(slide, x, y, w, 0.28, fill=color, line=color, lw=0.5, transparency=82)
    add_text(slide, label, x, y + 0.01, w, 0.22, size=size, color=color, bold=True,
             align=PP_ALIGN.CENTER, font=FONT_EN)
    return shape


def add_bullet(slide, content, x, y, w, color=BLUE, size=11):
    add_circle(slide, x, y + 0.105, 0.11, fill=color, line=color, lw=0.2)
    add_text(slide, content, x + 0.22, y, w - 0.22, 0.30, size=size, color=TEXT)


def add_placeholder(slide, x, y, w, h, label, note="后续替换图片", accent=BLUE):
    shape = box(slide, x, y, w, h, fill=BG, line=accent, lw=1.0, radius=False, dash=MSO_LINE_DASH_STYLE.DASH)
    shape.name = f"图片占位 - {label}"
    # Simple editable landscape icon.
    icon_x, icon_y = x + w / 2 - 0.33, y + h / 2 - 0.43
    outer = box(slide, icon_x, icon_y, 0.66, 0.48, fill=WHITE, line=accent, lw=0.8, radius=False)
    add_circle(slide, icon_x + 0.09, icon_y + 0.09, 0.10, fill=accent, line=accent, lw=0.2)
    add_line(slide, icon_x + 0.08, icon_y + 0.40, icon_x + 0.28, icon_y + 0.23, accent, 0.8)
    add_line(slide, icon_x + 0.28, icon_y + 0.23, icon_x + 0.43, icon_y + 0.36, accent, 0.8)
    add_line(slide, icon_x + 0.43, icon_y + 0.36, icon_x + 0.57, icon_y + 0.20, accent, 0.8)
    add_text(slide, "图片占位符", x + 0.20, y + h / 2 + 0.16, w - 0.40, 0.26, size=12,
             color=accent, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, label, x + 0.20, y + h / 2 + 0.48, w - 0.40, 0.32, size=10,
             color=MUTED, align=PP_ALIGN.CENTER)
    add_text(slide, f"【{note}】", x + 0.20, y + h - 0.38, w - 0.40, 0.18, size=8,
             color=MUTED, align=PP_ALIGN.CENTER)


def remove_placeholder(slide, idx):
    for shape in list(slide.placeholders):
        if shape.placeholder_format.idx == idx:
            element = shape._element
            element.getparent().remove(element)
            return


def set_page_number(slide, page):
    for shape in slide.placeholders:
        if shape.placeholder_format.idx == 12 and hasattr(shape, "text_frame"):
            set_text(shape, f"{page:02d}", size=10, color=BLUE, bold=True, align=PP_ALIGN.RIGHT, font=FONT_EN)
            return
    add_text(slide, f"{page:02d}", 11.88, 7.10, 0.80, 0.20, size=9, color=BLUE,
             bold=True, align=PP_ALIGN.RIGHT, font=FONT_EN)


def content_slide(prs, page, title, kicker):
    # Layout 14 is the template's "single image and text" content page.
    slide = prs.slides.add_slide(prs.slide_layouts[14])
    if slide.shapes.title is not None:
        set_text(slide.shapes.title, title, size=25, color=NAVY, bold=True)
    remove_placeholder(slide, 10)
    set_page_number(slide, page)
    add_text(slide, kicker.upper(), 0.40, 0.85, 3.10, 0.18, size=8, color=BLUE,
             bold=True, font=FONT_EN)
    add_line(slide, 0.40, 1.12, 12.50, 1.12, GRID, 0.7)
    return slide


def clear_slides(prs):
    slide_ids = prs.slides._sldIdLst
    for slide_id in list(slide_ids):
        prs.part.drop_rel(slide_id.rId)
        slide_ids.remove(slide_id)


def add_cover(prs):
    # Layout 11 carries the supplied template's cover graphic and Jiao Tong visual treatment.
    slide = prs.slides.add_slide(prs.slide_layouts[11])
    placeholders = {shape.placeholder_format.idx: shape for shape in slide.placeholders}
    set_text(placeholders[13], "项目介绍", size=15, color=BLUE, bold=True, font=FONT_EN)
    set_text(placeholders[0], "信息系统渗透\n智能化测试平台", size=29, color=NAVY, bold=True)
    set_text(placeholders[16], "面向信息系统安全测试的智能化平台\n建设目标 · 软硬件架构 · 核心能力 · 交付价值", size=14, color=MUTED)
    add_chip(slide, "INTELLIGENT", 5.66, 5.20, 1.35, CYAN, 8)
    add_chip(slide, "NATIVE", 7.12, 5.20, 1.02, GREEN, 8)
    add_chip(slide, "TRACEABLE", 8.25, 5.20, 1.30, PURPLE, 8)
    add_text(slide, "面向甲方项目介绍", 5.66, 5.74, 3.20, 0.20, size=10, color=MUTED, font=FONT_EN)
    return slide


def add_goal_solution(prs):
    slide = content_slide(prs, 2, "项目建设目标与总体方案", "01 / POSITIONING")
    add_text(slide, "以统一平台组织测试任务、工具资源与结果交付，形成安全测试业务闭环。", 0.58, 1.34, 7.25, 0.26, size=12, color=MUTED)
    box(slide, 0.58, 1.84, 6.98, 3.68, fill=WHITE, line=GRID)
    add_text(slide, "建设目标", 0.88, 2.08, 1.32, 0.28, size=15, color=NAVY, bold=True)
    goals = [
        "统一管理测试目标、测试任务和测试资源",
        "提升安全测试的自动化与智能化水平",
        "形成扫描发现到验证报告的完整闭环",
        "适配国产化软硬件与实验室、现场应用",
        "实现过程留痕、证据关联和结果追溯",
    ]
    for index, goal in enumerate(goals):
        add_bullet(slide, goal, 0.94, 2.60 + index * 0.50, 5.90, [BLUE, CYAN, GREEN, AMBER, PURPLE][index])
    add_text(slide, "总体方案", 0.88, 4.96, 1.32, 0.25, size=12, color=BLUE, bold=True)
    stages = [("人员", "测试 / 管理 / 教学", BLUE), ("客户端", "Qt/C++", CYAN), ("中心平台", "任务与协同", PURPLE), ("资源底座", "工具 / 模型 / 数据", GREEN)]
    for index, (label, detail, color) in enumerate(stages):
        x = 0.90 + index * 1.55
        box(slide, x, 5.32, 1.25, 0.62, fill=WHITE, line=color, lw=0.8)
        add_text(slide, label, x + 0.08, 5.40, 1.09, 0.16, size=9, color=color, bold=True, align=PP_ALIGN.CENTER)
        add_text(slide, detail, x + 0.08, 5.62, 1.09, 0.13, size=7.5, color=MUTED, align=PP_ALIGN.CENTER)
        if index < len(stages) - 1:
            add_chevron(slide, x + 1.30, 5.54, color=color)
    add_placeholder(slide, 8.05, 1.72, 4.18, 4.22, "平台整体示意图或应用场景图", accent=CYAN)
    add_text(slide, "统一入口 · 统一任务 · 统一资源 · 统一执行 · 统一报告", 0.80, 6.30, 11.58, 0.20, size=10, color=BLUE, bold=True, align=PP_ALIGN.CENTER)


def add_architecture(prs):
    slide = content_slide(prs, 3, "软硬件一体化架构", "02 / ARCHITECTURE")
    add_text(slide, "以中心服务器集中承载平台能力，以国产化云笔电作为多角色操作终端。", 0.58, 1.34, 9.30, 0.26, size=12, color=MUTED)
    box(slide, 0.58, 1.76, 5.42, 4.82, fill=WHITE, line=BLUE)
    add_text(slide, "硬件组成", 0.86, 2.00, 1.40, 0.26, size=15, color=NAVY, bold=True)
    hardware = [
        ("中心服务器", "海光 x86｜后端、数据库、容器、安全工具与模型服务", BLUE),
        ("云笔电客户端", "飞腾 E2000Q ARM64｜银河麒麟 V10 SP1｜Qt/C++", CYAN),
        ("测试网络", "PoE 局域网连接云笔电、测试靶机和被测设备", GREEN),
        ("扩展硬件", "GPU 加速、USB 硬件加密锁、虚拟机与网络设备", AMBER),
    ]
    for index, (title, detail, color) in enumerate(hardware):
        y = 2.48 + index * 0.91
        add_circle(slide, 0.90, y + 0.06, 0.28, fill=color, line=color, lw=0.2)
        add_text(slide, str(index + 1), 0.90, y + 0.12, 0.28, 0.10, size=7, color=WHITE, bold=True, align=PP_ALIGN.CENTER, font=FONT_EN)
        add_text(slide, title, 1.34, y, 1.24, 0.19, size=11, color=TEXT, bold=True)
        add_text(slide, detail, 1.34, y + 0.28, 4.20, 0.30, size=9, color=MUTED)
    box(slide, 6.34, 1.76, 5.90, 2.82, fill=WHITE, line=PURPLE)
    add_text(slide, "软件分层", 6.64, 2.00, 1.35, 0.26, size=15, color=NAVY, bold=True)
    layers = [("交互层", "Qt5/C++ 客户端", CYAN), ("服务层", "Node.js、任务调度、权限、WebSocket", BLUE), ("能力层", "扫描、分析、预案、执行、评分、报告", PURPLE), ("资源层", "SQLite、容器工具、知识库、模型、靶场", GREEN)]
    for index, (name, detail, color) in enumerate(layers):
        y = 2.45 + index * 0.49
        box(slide, 6.68, y, 1.14, 0.32, fill=color, line=color, lw=0.4, transparency=84)
        add_text(slide, name, 6.70, y + 0.06, 1.10, 0.14, size=8, color=color, bold=True, align=PP_ALIGN.CENTER)
        add_text(slide, detail, 8.04, y + 0.05, 3.72, 0.16, size=9, color=TEXT)
    add_placeholder(slide, 6.34, 4.88, 5.90, 1.70, "中心服务器、云笔电或部署现场照片", accent=GREEN)


def add_deployment(prs):
    slide = content_slide(prs, 4, "部署拓扑与运行模式", "03 / DEPLOYMENT")
    add_text(slide, "采用“中心服务 + 多端接入 + 测试环境”的集中部署方式，支持实验室与现场测试场景。", 0.58, 1.34, 10.20, 0.26, size=12, color=MUTED)
    box(slide, 0.58, 1.78, 8.10, 4.78, fill=WHITE, line=GRID)
    server = box(slide, 3.08, 2.08, 3.15, 0.78, fill=LIGHT_BLUE, line=BLUE, lw=1.0)
    server.name = "中心服务器"
    add_text(slide, "中心服务器｜海光 x86", 3.20, 2.22, 2.91, 0.20, size=13, color=BLUE, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "后端 · 数据库 · 容器 · 模型服务", 3.20, 2.50, 2.91, 0.16, size=8.5, color=MUTED, align=PP_ALIGN.CENTER)
    add_line(slide, 4.65, 2.86, 4.65, 3.38, BLUE, 1.2)
    add_text(slide, "PoE 局域网", 3.90, 3.05, 1.50, 0.18, size=10, color=BLUE, bold=True, align=PP_ALIGN.CENTER)
    add_line(slide, 1.30, 3.48, 7.92, 3.48, BLUE, 1.0)
    clients = [(0.90, "云笔电 1", "操作端", CYAN), (3.10, "云笔电 2", "控制端", PURPLE), (5.30, "云笔电 3~8", "扩展客户端", GREEN)]
    for x, name, role, color in clients:
        add_line(slide, x + 0.78, 3.48, x + 0.78, 3.88, color, 0.9)
        box(slide, x, 3.88, 1.56, 0.72, fill=WHITE, line=color, lw=0.8)
        add_text(slide, name, x + 0.08, 4.02, 1.40, 0.16, size=10, color=TEXT, bold=True, align=PP_ALIGN.CENTER)
        add_text(slide, role, x + 0.08, 4.27, 1.40, 0.14, size=8, color=color, align=PP_ALIGN.CENTER)
    add_line(slide, 4.65, 4.60, 4.65, 5.04, GREEN, 1.0)
    box(slide, 3.30, 5.04, 2.70, 0.66, fill=LIGHT_GREEN, line=GREEN, lw=0.8)
    add_text(slide, "测试靶场 / 被测系统", 3.46, 5.23, 2.38, 0.17, size=11, color=GREEN, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "任务与数据经服务器统一转发；云笔电之间无需直接通信。", 0.90, 5.96, 7.36, 0.20, size=9, color=MUTED, align=PP_ALIGN.CENTER)
    add_placeholder(slide, 9.04, 1.92, 3.20, 4.64, "网络部署拓扑图或现场连接照片", accent=CYAN)


def add_workflow(prs):
    slide = content_slide(prs, 5, "核心业务流程", "04 / WORKFLOW")
    add_text(slide, "围绕正式测试任务，形成从授权范围确认到测试成果交付的闭环管理。", 0.58, 1.34, 10.30, 0.26, size=12, color=MUTED)
    steps = [
        ("01", "任务准备", "目标 / 授权 / 策略", BLUE),
        ("02", "扫描发现", "资产 / 服务 / 风险线索", CYAN),
        ("03", "智能分析", "证据 / 知识 / 建议", PURPLE),
        ("04", "预案制定", "生成 / 编辑 / 审核", AMBER),
        ("05", "执行验证", "受控执行 / 证据采集", CORAL),
        ("06", "结果评估", "核验 / 报告 / 复测", GREEN),
    ]
    for index, (num, title, detail, color) in enumerate(steps):
        row = 0 if index < 3 else 1
        col = index if index < 3 else index - 3
        x = 0.72 + col * 2.56
        y = 1.96 + row * 1.60
        box(slide, x, y, 2.14, 1.02, fill=WHITE, line=color, lw=0.9)
        add_circle(slide, x + 0.16, y + 0.19, 0.34, fill=color, line=color, lw=0.2)
        add_text(slide, num, x + 0.16, y + 0.29, 0.34, 0.10, size=7, color=WHITE, bold=True, align=PP_ALIGN.CENTER, font=FONT_EN)
        add_text(slide, title, x + 0.66, y + 0.18, 1.28, 0.18, size=11, color=TEXT, bold=True)
        add_text(slide, detail, x + 0.18, y + 0.61, 1.78, 0.18, size=8.5, color=MUTED, align=PP_ALIGN.CENTER)
        if col < 2:
            add_chevron(slide, x + 2.22, y + 0.42, color=color)
    add_line(slide, 3.93, 3.08, 3.93, 3.38, PURPLE, 1.0)
    add_chevron(slide, 3.80, 3.22, 0.28, 0.20, PURPLE)
    box(slide, 0.72, 5.38, 7.36, 0.62, fill=LIGHT_PURPLE, line=PURPLE, lw=0.6)
    add_text(slide, "任务上下文贯穿全过程：补充扫描、重新分析、预案修订、重复验证和结果复核均可持续关联。", 0.98, 5.58, 6.84, 0.18, size=9.5, color=PURPLE, bold=True, align=PP_ALIGN.CENTER)
    add_placeholder(slide, 8.68, 1.96, 3.56, 4.04, "任务流程或任务详情界面截图", accent=AMBER)


def add_modules(prs):
    slide = content_slide(prs, 6, "平台功能模块与资源体系", "05 / CAPABILITIES")
    add_text(slide, "七个一级模块覆盖测试组织、专业分析、资源沉淀、成果管理和运行保障。", 0.58, 1.34, 10.00, 0.26, size=12, color=MUTED)
    modules = [
        ("工作台", "全局态势、待办与运行状态", BLUE),
        ("测试任务", "任务创建、目标配置与流程管理", CYAN),
        ("资产与拓扑", "资产、服务、网络关系展示", GREEN),
        ("智能分析", "风险识别、证据关联与建议", PURPLE),
        ("预案与知识库", "Playbook、知识图谱、漏洞与工具", AMBER),
        ("结果与报告", "验证结果、证据管理与报告导出", CORAL),
        ("系统管理", "用户、权限、模型、工具与审计", BLUE),
    ]
    for index, (name, detail, color) in enumerate(modules):
        x = 0.70 + (index % 2) * 3.56
        y = 1.88 + (index // 2) * 0.90
        box(slide, x, y, 3.22, 0.66, fill=WHITE, line=color, lw=0.7)
        add_text(slide, name, x + 0.17, y + 0.14, 1.04, 0.16, size=10.5, color=color, bold=True)
        add_text(slide, detail, x + 1.24, y + 0.14, 1.80, 0.28, size=8.5, color=MUTED)
    box(slide, 0.70, 5.60, 6.78, 0.62, fill=LIGHT_BLUE, line=BLUE, lw=0.6)
    add_text(slide, "资源体系：网络探测 · 漏洞扫描 · Web 安全 · 凭证测试 · 漏洞验证 · 云安全 · 系统检测", 0.95, 5.78, 6.28, 0.16, size=9, color=BLUE, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "工具通过 Podman/Docker 容器统一管理，便于部署、迁移和版本维护。", 0.74, 6.36, 6.75, 0.16, size=8.8, color=MUTED, align=PP_ALIGN.CENTER)
    add_placeholder(slide, 8.10, 1.86, 4.14, 4.36, "工作台、资产拓扑或功能模块界面截图", accent=GREEN)


def add_ai(prs):
    slide = content_slide(prs, 7, "AI 智能协作能力", "06 / AI COLLABORATION")
    add_text(slide, "AI 作为任务协作助手，连接扫描事实、知识资源、结构化预案和测试报告。", 0.58, 1.34, 9.90, 0.26, size=12, color=MUTED)
    box(slide, 0.62, 1.86, 7.05, 4.48, fill=WHITE, line=PURPLE)
    nodes = [(1.08, 2.30, "扫描事实", "资产、端口、服务、发现", CYAN), (3.25, 2.30, "智能分析", "风险候选、证据、建议", PURPLE), (5.42, 2.30, "预案与报告", "方案生成、结果整理", GREEN)]
    for index, (x, y, title, detail, color) in enumerate(nodes):
        add_circle(slide, x, y, 1.28, fill=WHITE, line=color, lw=1.3)
        add_text(slide, str(index + 1), x + 0.48, y + 0.22, 0.32, 0.18, size=10, color=color, bold=True, align=PP_ALIGN.CENTER, font=FONT_EN)
        add_text(slide, title, x + 0.08, y + 0.53, 1.12, 0.18, size=10, color=TEXT, bold=True, align=PP_ALIGN.CENTER)
        add_text(slide, detail, x - 0.26, y + 1.42, 1.80, 0.32, size=8.2, color=MUTED, align=PP_ALIGN.CENTER)
        if index < 2:
            add_chevron(slide, x + 1.42, y + 0.53, 0.42, 0.22, color)
    box(slide, 1.02, 4.58, 6.24, 1.10, fill=LIGHT_PURPLE, line=PURPLE, lw=0.7)
    add_text(slide, "AI 协作模式", 1.28, 4.82, 1.20, 0.18, size=11, color=PURPLE, bold=True)
    add_text(slide, "AI 辅助决策  →  人工审核确认  →  系统受控执行", 2.52, 4.80, 4.24, 0.22, size=11, color=TEXT, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "支持自然语言问答、依据解释、预案修订、证据整理和报告辅助编写。", 1.28, 5.22, 5.74, 0.16, size=9, color=MUTED, align=PP_ALIGN.CENTER)
    add_placeholder(slide, 8.06, 1.86, 4.18, 4.48, "AI 协作、预案编辑或智能分析界面截图", accent=PURPLE)


def add_security(prs):
    slide = content_slide(prs, 8, "权限、安全与审计保障", "07 / SECURITY & AUDIT")
    add_text(slide, "以授权边界、人工审核、执行留痕和证据追溯保障测试过程安全可控。", 0.58, 1.34, 10.10, 0.26, size=12, color=MUTED)
    box(slide, 0.64, 1.82, 7.06, 4.52, fill=WHITE, line=GREEN)
    add_text(slide, "安全控制链", 0.94, 2.08, 1.32, 0.25, size=15, color=NAVY, bold=True)
    controls = [
        ("身份与权限", "用户、角色和权限分级管理", BLUE),
        ("授权与审批", "测试范围、执行边界和高风险操作控制", AMBER),
        ("过程留痕", "工具调用、日志、原始输出与状态记录", CYAN),
        ("证据追溯", "任务、预案、执行记录与报告相互关联", GREEN),
    ]
    for index, (title, detail, color) in enumerate(controls):
        y = 2.56 + index * 0.74
        add_circle(slide, 1.00, y + 0.02, 0.28, fill=color, line=color, lw=0.2)
        add_text(slide, "✓", 1.00, y + 0.09, 0.28, 0.10, size=8, color=WHITE, bold=True, align=PP_ALIGN.CENTER)
        add_text(slide, title, 1.46, y, 1.32, 0.18, size=10.5, color=TEXT, bold=True)
        add_text(slide, detail, 2.90, y, 4.18, 0.30, size=9, color=MUTED)
    box(slide, 0.96, 5.66, 6.34, 0.38, fill=LIGHT_GREEN, line=GREEN, lw=0.5)
    add_text(slide, "可结合 USB 硬件加密锁实现软件授权控制。", 1.20, 5.76, 5.86, 0.14, size=9, color=GREEN, bold=True, align=PP_ALIGN.CENTER)
    add_placeholder(slide, 8.06, 1.82, 4.18, 4.52, "权限管理、审批或审计日志界面截图", accent=GREEN)


def add_results(prs):
    slide = content_slide(prs, 9, "结果评估与报告交付", "08 / RESULTS & REPORT")
    add_text(slide, "将测试过程转化为可审核、可追溯、可交付的结构化成果。", 0.58, 1.34, 9.30, 0.26, size=12, color=MUTED)
    add_placeholder(slide, 0.62, 1.84, 4.54, 4.56, "测试报告封面、结果总览或报告页面截图", accent=CORAL)
    box(slide, 5.56, 1.84, 6.68, 4.56, fill=WHITE, line=CORAL)
    add_text(slide, "甲方可获得的测试成果", 5.88, 2.10, 2.68, 0.26, size=15, color=NAVY, bold=True)
    deliverables = [
        "测试任务总体概览与资产、服务清单",
        "风险发现及经过验证的测试结果",
        "攻击步骤、工具执行记录与网络数据",
        "原始输出、日志、证据和风险评分",
        "结构化测试报告、审核、导出和归档文件",
    ]
    for index, item in enumerate(deliverables):
        add_bullet(slide, item, 5.94, 2.62 + index * 0.54, 5.72, [BLUE, CYAN, GREEN, AMBER, CORAL][index], 10.5)
    box(slide, 5.88, 5.60, 5.98, 0.46, fill=LIGHT_AMBER, line=AMBER, lw=0.6)
    add_text(slide, "报告 → 任务 → 预案 → 执行记录 → 原始证据", 6.10, 5.74, 5.54, 0.14, size=9, color=AMBER, bold=True, align=PP_ALIGN.CENTER)


def add_scenarios(prs):
    slide = content_slide(prs, 10, "典型应用场景与应用价值", "09 / SCENARIOS & VALUE")
    add_text(slide, "面向实验室、教学实训、国产化适配和现场检测等多类安全测试应用。", 0.58, 1.34, 10.00, 0.26, size=12, color=MUTED)
    scenarios = [
        ("实验室测试", "网络安全测试与环境验证", BLUE),
        ("国产化适配", "软硬件环境兼容性测试", CYAN),
        ("云笔电协同", "多角色任务协同与控制", PURPLE),
        ("靶场实训", "教学、演练与能力评估", GREEN),
        ("现场检测", "授权范围内的安全检测", AMBER),
        ("结果复盘", "报告归档与安全改进", CORAL),
    ]
    for index, (title, detail, color) in enumerate(scenarios):
        x = 0.70 + (index % 2) * 3.34
        y = 1.92 + (index // 2) * 1.04
        box(slide, x, y, 2.98, 0.78, fill=WHITE, line=color, lw=0.7)
        add_circle(slide, x + 0.16, y + 0.18, 0.36, fill=color, line=color, lw=0.2)
        add_text(slide, str(index + 1), x + 0.16, y + 0.30, 0.36, 0.10, size=7, color=WHITE, bold=True, align=PP_ALIGN.CENTER, font=FONT_EN)
        add_text(slide, title, x + 0.68, y + 0.15, 1.90, 0.17, size=10.5, color=TEXT, bold=True)
        add_text(slide, detail, x + 0.68, y + 0.42, 2.08, 0.14, size=8.2, color=MUTED)
    box(slide, 0.70, 5.32, 6.64, 0.82, fill=LIGHT_BLUE, line=BLUE, lw=0.6)
    add_text(slide, "应用价值", 0.98, 5.56, 0.94, 0.18, size=11, color=BLUE, bold=True)
    add_text(slide, "提升测试效率 · 降低工具使用门槛 · 规范测试流程 · 强化结果可信度与可追溯性", 2.02, 5.55, 5.02, 0.20, size=9.2, color=TEXT, bold=True, align=PP_ALIGN.CENTER)
    add_placeholder(slide, 8.00, 1.92, 4.24, 4.22, "实验室、云笔电协同或现场测试照片", accent=BLUE)


def add_delivery(prs):
    slide = content_slide(prs, 11, "部署运维、扩展与交付内容", "10 / DELIVERY")
    add_text(slide, "平台支持集中部署与持续扩展，并形成软件、资源、文档和环境配置等交付物。", 0.58, 1.34, 10.40, 0.26, size=12, color=MUTED)
    box(slide, 0.62, 1.84, 5.34, 3.58, fill=WHITE, line=BLUE)
    add_text(slide, "部署与扩展能力", 0.92, 2.10, 1.78, 0.24, size=15, color=NAVY, bold=True)
    ops = [
        "服务器集中部署，客户端批量接入",
        "支持离线安装、物理迁移和工具镜像更新",
        "适配国产操作系统与多种终端环境",
        "可扩展安全工具、模型、测试靶场和专用硬件",
    ]
    for index, item in enumerate(ops):
        add_bullet(slide, item, 0.96, 2.60 + index * 0.53, 4.58, [BLUE, CYAN, GREEN, PURPLE][index], 10)
    box(slide, 6.34, 1.84, 5.90, 3.58, fill=WHITE, line=GREEN)
    add_text(slide, "项目交付内容", 6.64, 2.10, 1.50, 0.24, size=15, color=NAVY, bold=True)
    items = [
        "平台客户端与中心服务端软件",
        "数据库、数据管理模块与安全工具容器镜像",
        "AI 模型接入、测试任务、报告与靶场接入能力",
        "部署手册、用户手册、运维资料与硬件网络说明",
    ]
    for index, item in enumerate(items):
        add_bullet(slide, item, 6.68, 2.60 + index * 0.53, 5.04, [GREEN, BLUE, AMBER, CYAN][index], 10)
    add_placeholder(slide, 0.62, 5.74, 11.62, 0.70, "硬件设备、软件安装包或交付清单示意图", note="可替换为交付设备与软件目录照片", accent=AMBER)


def add_closing(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[16])
    # The supplied tail-page layout is preserved; custom text provides the project close.
    add_text(slide, "信息系统渗透智能化测试平台", 2.05, 2.06, 9.24, 0.36, size=17, color=BLUE, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "构建统一、智能、可追溯的安全测试工作平台", 1.55, 2.62, 10.24, 0.76, size=29, color=NAVY, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "软硬件协同部署 · 全流程任务闭环 · 人机协作分析 · 结果证据追溯", 1.55, 3.55, 10.24, 0.26, size=13, color=MUTED, align=PP_ALIGN.CENTER)
    add_line(slide, 4.38, 4.24, 8.96, 4.24, BLUE, 1.0)
    add_text(slide, "谢谢各位领导、老师，欢迎指导", 1.55, 4.66, 10.24, 0.34, size=17, color=TEXT, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "THANK YOU", 1.55, 5.26, 10.24, 0.22, size=10, color=BLUE, bold=True, align=PP_ALIGN.CENTER, font=FONT_EN)


def build():
    prs = Presentation(str(TEMPLATE))
    clear_slides(prs)
    add_cover(prs)
    add_goal_solution(prs)
    add_architecture(prs)
    add_deployment(prs)
    add_workflow(prs)
    add_modules(prs)
    add_ai(prs)
    add_security(prs)
    add_results(prs)
    add_scenarios(prs)
    add_delivery(prs)
    add_closing(prs)
    prs.save(str(OUTPUT))
    print(f"created: {OUTPUT}")
    print(f"slides: {len(prs.slides)}")


if __name__ == "__main__":
    build()
