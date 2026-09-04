from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.dml.color import RGBColor


TEMPLATE = Path(__file__).resolve().parents[1] / "docs" / "软件需规" / "交大ppt模板.pptx"
OUT = Path(__file__).resolve().parents[1] / "docs" / "软件需规" / "RedTeam-Platform_8.30以来专项工作汇报_交大模板.pptx"

W, H = 13.333, 7.5
WHITE = RGBColor(255, 255, 255)
BG = RGBColor(249, 251, 254)
NAVY = RGBColor(15, 20, 35)
BLUE = RGBColor(96, 150, 230)      # template accent1
PURPLE = RGBColor(149, 77, 114)    # template accent2-ish supporting tone
CYAN = RGBColor(88, 182, 229)      # template accent2
GREEN = RGBColor(86, 202, 149)     # template accent3
AMBER = RGBColor(255, 186, 85)     # template accent4
CORAL = RGBColor(241, 136, 112)    # template accent5
RED = RGBColor(236, 95, 116)       # template accent6
TEXT = RGBColor(32, 48, 72)
MUTED = RGBColor(96, 111, 132)
GRID = RGBColor(222, 230, 240)
LIGHT_BLUE = RGBColor(241, 247, 253)
LIGHT_PURPLE = RGBColor(246, 243, 253)
LIGHT_GREEN = RGBColor(241, 251, 247)
LIGHT_AMBER = RGBColor(255, 249, 239)
FONT = "Microsoft YaHei"
MONO = "Arial"


def fill_shape(shape, color, transparency=0):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.fill.transparency = transparency


def line_shape(shape, color, width=1.0, transparency=0):
    shape.line.color.rgb = color
    shape.line.width = Pt(width)
    shape.line.transparency = transparency


def rect(slide, x, y, w, h, fill=WHITE, line=GRID, lw=0.7, radius=True, transparency=0):
    typ = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shp = slide.shapes.add_shape(typ, Inches(x), Inches(y), Inches(w), Inches(h))
    fill_shape(shp, fill, transparency)
    line_shape(shp, line, lw)
    return shp


def text(slide, content, x, y, w, h, size=12, color=TEXT, bold=False, align=PP_ALIGN.LEFT,
         font=FONT, margin=0.04, valign=MSO_ANCHOR.MIDDLE, italic=False):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear(); tf.word_wrap = True
    tf.margin_left = Inches(margin); tf.margin_right = Inches(margin)
    tf.margin_top = Inches(margin); tf.margin_bottom = Inches(margin)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]; p.alignment = align
    run = p.add_run(); run.text = content
    run.font.name = font; run.font.size = Pt(size); run.font.bold = bold
    run.font.italic = italic; run.font.color.rgb = color
    return box


def line(slide, x1, y1, x2, y2, color=GRID, width=1.0, transparency=0):
    ln = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    line_shape(ln, color, width, transparency)
    return ln


def arrow(slide, x, y, w, h, color=BLUE, transparency=0):
    shp = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(x), Inches(y), Inches(w), Inches(h))
    fill_shape(shp, color, transparency); line_shape(shp, color, 0.1, 100)
    return shp


def circle(slide, x, y, d, fill=WHITE, line_color=BLUE, lw=1.0, transparency=0):
    shp = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d))
    fill_shape(shp, fill, transparency); line_shape(shp, line_color, lw)
    return shp


def chip(slide, label, x, y, w, color=BLUE, fill=None, size=8):
    fill = fill or color
    shp = rect(slide, x, y, w, 0.28, fill=fill, line=color, lw=0.6, transparency=78)
    text(slide, label, x, y + 0.01, w, 0.24, size=size, color=color, bold=True, align=PP_ALIGN.CENTER, font=MONO)
    return shp


def set_shape_text(shape, content, size=None, color=None, bold=None, align=None):
    if not hasattr(shape, "text_frame"):
        return
    tf = shape.text_frame
    tf.clear(); tf.word_wrap = True
    p = tf.paragraphs[0]
    if align is not None: p.alignment = align
    r = p.add_run(); r.text = content
    if size is not None: r.font.size = Pt(size)
    if color is not None: r.font.color.rgb = color
    if bold is not None: r.font.bold = bold
    r.font.name = FONT


def section_slide(slide, numeral, title, subtitle):
    # Existing divider slide already carries the template's visual language.
    if len(slide.shapes) >= 5:
        set_shape_text(slide.shapes[0], numeral, size=56, color=BLUE, bold=True, align=PP_ALIGN.CENTER)
        set_shape_text(slide.shapes[1], title + "  ", size=26, color=NAVY, bold=True)
        set_shape_text(slide.shapes[4], subtitle, size=14, color=MUTED)


def add_content_title(slide, title, kicker):
    text(slide, kicker.upper(), 0.70, 0.24, 2.30, 0.18, size=9, color=BLUE, bold=True, font=MONO)
    text(slide, title, 0.70, 0.55, 10.60, 0.42, size=22, color=NAVY, bold=True)
    line(slide, 0.70, 1.09, 12.58, 1.09, color=GRID, width=0.8)


def add_footer_note(slide, note, color=BLUE):
    rect(slide, 0.72, 6.44, 11.80, 0.34, fill=BG, line=GRID, lw=0.5)
    text(slide, note, 0.92, 6.51, 11.40, 0.18, size=9, color=color, bold=True, align=PP_ALIGN.CENTER)


def build():
    prs = Presentation(str(TEMPLATE))
    # Template structure: cover, agenda, 4 section dividers + 4 content slides, thanks.

    # 1 Cover
    slide = prs.slides[0]
    # The template's patterned strip remains as a visual anchor; cover text is added above it.
    text(slide, "阶段工作汇报", 0.74, 1.10, 4.80, 0.58, size=30, color=NAVY, bold=True)
    text(slide, "2026.08.30 — 至今", 0.78, 1.83, 3.50, 0.24, size=12, color=BLUE, bold=True, font=MONO)
    text(slide, "信息系统渗透智能化测试平台", 0.78, 2.46, 5.70, 0.32, size=17, color=TEXT, bold=True)
    text(slide, "RedTeam-Platform", 0.78, 2.88, 4.80, 0.24, size=14, color=MUTED, font=MONO)
    text(slide, "本阶段围绕攻击编排、扫描优化、评估输出和系统演示，完成专项功能补强。", 0.78, 3.50, 5.58, 0.52, size=13, color=MUTED)
    for i, (lab, col) in enumerate([("网络捕获", CYAN), ("扫描优化", BLUE), ("三类攻击", RED), ("智能报告", GREEN)]):
        chip(slide, lab, 0.78 + i * 1.40, 4.52, 1.14, col, size=8)
    # Replace template's small word art label with project name marker.
    if len(slide.shapes) > 1 and hasattr(slide.shapes[1], "text_frame"):
        set_shape_text(slide.shapes[1], "RedTeam-Platform", size=10, color=BLUE, bold=True)
    text(slide, "PROJECT UPDATE  ·  FOR SYSTEM DEMO", 8.10, 6.58, 4.06, 0.18, size=8, color=MUTED, align=PP_ALIGN.RIGHT, font=MONO)

    # 2 Agenda
    slide = prs.slides[1]
    # Clear the template placeholder text but preserve its background.
    if len(slide.shapes) > 0 and hasattr(slide.shapes[0], "text_frame"):
        set_shape_text(slide.shapes[0], "")
    text(slide, "本次汇报内容", 0.78, 0.84, 3.40, 0.46, size=25, color=NAVY, bold=True)
    text(slide, "四项近期工作，对应四段系统演示主线", 0.80, 1.38, 5.40, 0.24, size=12, color=MUTED)
    agenda = [("01", "网络数据捕获", "采集任务 · PCAP · 协议分析", CYAN), ("02", "脆弱性扫描优化", "应用识别 · 目标分类 · 推荐", BLUE), ("03", "三类攻击实现", "Campaign · 阶段 · 产物传递", RED), ("04", "测试报告智能生成", "证据汇总 · 预览 · 多格式导出", GREEN)]
    for i, (num, title, desc, col) in enumerate(agenda):
        x = 0.84 + (i % 2) * 6.06; y = 2.20 + (i // 2) * 1.62
        rect(slide, x, y, 5.48, 1.22, fill=WHITE, line=col, lw=0.8)
        circle(slide, x + 0.22, y + 0.32, 0.42, fill=col, line_color=col, lw=0.4)
        text(slide, num, x + 0.22, y + 0.45, 0.42, 0.14, size=9, color=BG, bold=True, align=PP_ALIGN.CENTER, font=MONO)
        text(slide, title, x + 0.90, y + 0.22, 2.30, 0.22, size=13, color=TEXT, bold=True)
        text(slide, desc, x + 0.90, y + 0.62, 3.72, 0.22, size=10, color=MUTED)
    add_footer_note(slide, "演示顺序：先展示发现，再展示攻击决策，最后落到证据与报告。", CYAN)

    # 3 / 4 Network capture
    section_slide(prs.slides[2], "一", "网络数据捕获", "从流量采集到分析证据")
    slide = prs.slides[3]; add_content_title(slide, "网络数据捕获：把攻击过程留下来", "01 / CAPTURE")
    text(slide, "新增捕获任务、启动/停止控制、PCAP 文件管理与 tshark 分析服务，接入测试评估页面。", 0.72, 1.30, 10.40, 0.28, size=12, color=MUTED)
    # Flow diagram
    steps = [("01", "创建任务", "接口 / 网卡 / 过滤条件", CYAN), ("02", "开始捕获", "tcpdump 容器执行", BLUE), ("03", "停止并落盘", "生成 PCAP 文件", PURPLE), ("04", "分析结果", "协议 / 会话 / 异常", GREEN)]
    for i, (num, title, desc, col) in enumerate(steps):
        x = 0.84 + i * 3.02
        rect(slide, x, 2.18, 2.36, 1.14, fill=WHITE, line=col, lw=0.8)
        text(slide, num, x + 0.18, 2.38, 0.34, 0.18, size=10, color=col, bold=True, font=MONO)
        text(slide, title, x + 0.66, 2.34, 1.40, 0.22, size=12, color=TEXT, bold=True)
        text(slide, desc, x + 0.18, 2.78, 1.96, 0.26, size=9, color=MUTED)
        if i < 3: arrow(slide, x + 2.42, 2.66, 0.38, 0.16, col, transparency=10)
    # Lower detail cards
    for x, title, desc, col in [(0.84, "任务生命周期", "created → running → stopped → analyzed", CYAN), (4.02, "分析服务", "tshark 协议统计、会话与异常识别", PURPLE), (7.20, "前端入口", "评估页独立 Tab，状态实时刷新", GREEN), (10.38, "当前状态", "功能已接入，现场环境持续联调", AMBER)]:
        rect(slide, x, 4.20, 2.36, 1.26, fill=WHITE, line=col, lw=0.7)
        text(slide, title, x + 0.18, 4.46, 1.92, 0.20, size=11, color=TEXT, bold=True)
        text(slide, desc, x + 0.18, 4.86, 1.92, 0.36, size=9, color=MUTED)
    add_footer_note(slide, "演示建议：评估页 → 网络数据捕获 → 新建任务 → 开始/停止 → 查看分析结果。", CYAN)

    # 5 / 6 Vulnerability scanning
    section_slide(prs.slides[4], "二", "脆弱性扫描优化", "让扫描结果直接驱动下一步")
    slide = prs.slides[5]; add_content_title(slide, "脆弱性扫描优化：面向应用目标的发现与推荐", "02 / SCAN")
    text(slide, "扫描页面从单一任务列表扩展为“目标识别 + 多类扫描 + 结果分析 + Playbook 推荐”工作台。", 0.72, 1.30, 10.60, 0.28, size=12, color=MUTED)
    # Left: scan engines
    rect(slide, 0.78, 2.00, 4.12, 3.96, fill=WHITE, line=BLUE, lw=0.8)
    text(slide, "新增应用扫描类型", 1.08, 2.28, 2.20, 0.22, size=13, color=TEXT, bold=True)
    for i, (tool, desc, col) in enumerate([("whatweb + httpx", "技术栈与 HTTP 特征发现", CYAN), ("arjun + ffuf", "API 端点与参数模糊测试", PURPLE), ("nuclei 认证模板", "认证机制与配置审计", GREEN)]):
        y = 2.90 + i * 0.72
        circle(slide, 1.10, y + 0.04, 0.20, fill=col, line_color=col, lw=0.3)
        text(slide, tool, 1.48, y, 1.96, 0.20, size=10, color=TEXT, bold=True, font=MONO)
        text(slide, desc, 1.48, y + 0.28, 2.74, 0.20, size=9, color=MUTED)
    line(slide, 1.08, 5.10, 4.56, 5.10, color=GRID, width=0.6)
    text(slide, "目标分类", 1.08, 5.34, 0.88, 0.18, size=9, color=MUTED, bold=True)
    for i, lab in enumerate(["Web", "REST", "GraphQL", "SPA"]): chip(slide, lab, 2.08 + i * 0.62, 5.30, 0.54, [CYAN, PURPLE, AMBER, GREEN][i], size=7)
    # Right: result to playbook
    rect(slide, 5.18, 2.00, 7.42, 3.96, fill=WHITE, line=CYAN, lw=0.8)
    text(slide, "结果驱动的推荐链路", 5.48, 2.28, 2.50, 0.22, size=13, color=TEXT, bold=True)
    chain = [("发现", "443 / Spring / /api", CYAN), ("分析", "技术栈 + 攻击面", BLUE), ("匹配", "现有 Playbook", PURPLE), ("预览", "步骤与参数", GREEN)]
    for i, (title, desc, col) in enumerate(chain):
        x = 5.50 + i * 1.72
        circle(slide, x, 3.16, 0.42, fill=col, line_color=col, lw=0.4)
        text(slide, str(i + 1), x, 3.30, 0.42, 0.14, size=8, color=BG, bold=True, align=PP_ALIGN.CENTER, font=MONO)
        text(slide, title, x - 0.14, 3.78, 0.70, 0.18, size=10, color=TEXT, bold=True, align=PP_ALIGN.CENTER)
        text(slide, desc, x - 0.46, 4.10, 1.34, 0.32, size=8, color=MUTED, align=PP_ALIGN.CENTER)
        if i < 3: arrow(slide, x + 0.62, 3.32, 0.42, 0.16, col, transparency=10)
    rect(slide, 5.50, 5.18, 6.68, 0.48, fill=LIGHT_BLUE, line=GRID, lw=0.5)
    text(slide, "推荐项点击先查看；AI 生成预案先预览，确认后再前往执行。", 5.74, 5.32, 6.18, 0.18, size=9, color=CYAN, bold=True, align=PP_ALIGN.CENTER)
    add_footer_note(slide, "演示建议：创建应用目标 → 运行扫描 → 查看发现 → 展开推荐预案 → 前往执行。", BLUE)

    # 7 / 8 Attack orchestration
    section_slide(prs.slides[6], "三", "三类攻击实现", "从分类标签到阶段化战役")
    slide = prs.slides[7]; add_content_title(slide, "三类攻击实现：用 Campaign 组织攻击行动", "03 / ATTACK")
    text(slide, "新增 Campaign 数据模型、阶段管理、Playbook 编排、阶段产物传递与运行控制。", 0.72, 1.30, 10.30, 0.28, size=12, color=MUTED)
    phases = [("01", "数据抵近窃取", "凭据 / 文件 / 数据", CYAN), ("02", "信息篡改欺骗", "DNS / 会话 / 误导", PURPLE), ("03", "关键设备夺控", "横向 / 控制 / 持久", RED)]
    for i, (num, title, desc, col) in enumerate(phases):
        x = 0.84 + i * 3.04
        rect(slide, x, 2.10, 2.54, 1.24, fill=WHITE, line=col, lw=0.9)
        circle(slide, x + 0.20, 2.34, 0.36, fill=col, line_color=col, lw=0.3)
        text(slide, num, x + 0.20, 2.46, 0.36, 0.14, size=8, color=BG, bold=True, align=PP_ALIGN.CENTER, font=MONO)
        text(slide, title, x + 0.70, 2.32, 1.54, 0.22, size=11, color=TEXT, bold=True)
        text(slide, desc, x + 0.20, 2.86, 2.10, 0.20, size=9, color=MUTED)
        if i < 2: arrow(slide, x + 2.60, 2.62, 0.38, 0.16, col, transparency=10)
    # lower architecture
    rect(slide, 0.84, 3.88, 11.82, 2.02, fill=WHITE, line=GRID, lw=0.7)
    cols = [("阶段", "每个阶段可跳过、推进、查看状态", CYAN), ("预案", "阶段内支持串行 / 并行 / 条件执行", PURPLE), ("产物", "凭据、IP、Token 注入后续步骤", GREEN), ("控制", "启动 / 暂停 / 终止 / 战役报告", AMBER)]
    for i, (title, desc, col) in enumerate(cols):
        x = 1.12 + i * 2.86
        circle(slide, x, 4.36, 0.26, fill=col, line_color=col, lw=0.3)
        text(slide, title, x + 0.42, 4.32, 1.56, 0.20, size=11, color=TEXT, bold=True)
        text(slide, desc, x, 4.84, 2.24, 0.42, size=9, color=MUTED)
    add_footer_note(slide, "演示建议：新建战役 → 为三个阶段添加 Playbook → 查看阶段产物 → 启动战役。", RED)

    # 9 / 10 Reports
    section_slide(prs.slides[8], "四", "测试报告智能生成", "从执行记录到可交付报告")
    slide = prs.slides[9]; add_content_title(slide, "测试报告智能生成：让证据自动变成结论", "04 / REPORT")
    text(slide, "基于执行记录、攻击证据和捕获分析结果，自动生成结构化报告，并支持预览与导出。", 0.72, 1.30, 10.60, 0.28, size=12, color=MUTED)
    # report pipeline
    stages = [("证据汇总", "run / steps / evidence", CYAN), ("智能整理", "风险、影响、结论", PURPLE), ("报告预览", "JSON ↔ HTML", BLUE), ("交付导出", "DOCX / PDF / HTML", GREEN)]
    for i, (title, desc, col) in enumerate(stages):
        x = 0.84 + i * 3.02
        rect(slide, x, 2.18, 2.36, 1.12, fill=WHITE, line=col, lw=0.8)
        text(slide, title, x + 0.18, 2.44, 1.82, 0.20, size=12, color=TEXT, bold=True)
        text(slide, desc, x + 0.18, 2.82, 1.92, 0.18, size=9, color=MUTED, font=MONO)
        if i < 3: arrow(slide, x + 2.42, 2.66, 0.42, 0.16, col, transparency=10)
    details = [(0.84, "报告预览", "源码与 HTML 双模式", CYAN), (4.02, "多格式导出", "DOCX / PDF / HTML", GREEN), (7.20, "WPS 打开", "导出后可直接编辑", PURPLE), (10.38, "模板接口", "生成与导出 API", AMBER)]
    for x, title, desc, col in details:
        rect(slide, x, 4.14, 2.36, 1.30, fill=WHITE, line=col, lw=0.7)
        text(slide, title, x + 0.18, 4.40, 1.82, 0.20, size=11, color=TEXT, bold=True)
        text(slide, desc, x + 0.18, 4.80, 1.92, 0.22, size=9, color=MUTED)
    add_footer_note(slide, "演示建议：选择执行记录 → 生成测试报告 → HTML 预览 → 导出 DOCX/PDF。", GREEN)

    # 11 Closing
    slide = prs.slides[10]
    if len(slide.shapes) > 0 and hasattr(slide.shapes[0], "text_frame"):
        set_shape_text(slide.shapes[0], "感谢聆听", size=25, color=NAVY, bold=True)
    if len(slide.shapes) > 2 and hasattr(slide.shapes[2], "text_frame"):
        set_shape_text(slide.shapes[2], "RedTeam-Platform", size=16, color=BLUE, bold=True, align=PP_ALIGN.CENTER)
    if len(slide.shapes) > 3 and hasattr(slide.shapes[3], "text_frame"):
        set_shape_text(slide.shapes[3], "谢谢", size=36, color=NAVY, bold=True, align=PP_ALIGN.CENTER)
    text(slide, "系统演示主线：扫描 → 推荐 → 执行 → 证据 → 报告", 3.90, 5.18, 5.60, 0.24, size=12, color=BLUE, bold=True, align=PP_ALIGN.CENTER)

    return prs


if __name__ == "__main__":
    OUT.parent.mkdir(parents=True, exist_ok=True)
    build().save(str(OUT))
    print(OUT)
