from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor

# Reuse the light technical visual language and drawing helpers from the project-intro deck.
from generate_redteam_intro_ppt import (
    W, H, BG, PANEL, PANEL_2, TEXT, MUTED, CYAN, BLUE, PURPLE, GREEN, AMBER, RED, GRID,
    FONT, MONO, set_fill, set_line, add_rect, add_text, add_line, add_arrow, add_circle,
    add_chip, add_header, add_dark_bg, draw_shield,
)


def add_card_title(slide, title, x, y, w, color=TEXT):
    add_text(slide, title, x, y, w, 0.28, size=13, color=color, bold=True)


def add_status(slide, label, x, y, color=GREEN, width=0.72):
    add_chip(slide, label, x, y, width, color, size=8)


def add_timeline_node(slide, x, y, number, title, desc, color):
    add_circle(slide, x, y, 0.40, fill=color, line=color, lw=0.5)
    add_text(slide, number, x, y + 0.12, 0.40, 0.15, size=9, color=BG, bold=True, align=PP_ALIGN.CENTER, font=MONO)
    add_text(slide, title, x + 0.62, y + 0.01, 2.26, 0.22, size=12, color=TEXT, bold=True)
    add_text(slide, desc, x + 0.62, y + 0.30, 3.70, 0.36, size=10, color=MUTED)


def make_progress_deck():
    prs = Presentation()
    prs.slide_width = Inches(W)
    prs.slide_height = Inches(H)
    blank = prs.slide_layouts[6]

    # 1. Cover
    slide = prs.slides.add_slide(blank)
    add_dark_bg(slide, grid=False)
    add_text(slide, "阶段工作汇报", 0.78, 1.17, 3.6, 0.56, size=30, color=TEXT, bold=True)
    add_text(slide, "2026.07.29 — 2026.09.03", 0.82, 1.86, 3.6, 0.24, size=12, color=CYAN, bold=True, font=MONO)
    add_text(slide, "信息系统渗透智能化测试平台", 0.82, 2.52, 5.60, 0.34, size=17, color=TEXT, bold=True)
    add_text(slide, "RedTeam-Platform", 0.82, 2.95, 4.8, 0.28, size=15, color=MUTED, font=MONO)
    add_text(slide, "围绕会议要求，持续完善平台功能、攻击链路与部署验证能力。", 0.82, 3.57, 5.30, 0.44, size=13, color=MUTED)
    for i, (value, label, col) in enumerate([("30", "可执行工具", CYAN), ("7", "功能模块", PURPLE), ("4", "应用目标类型", GREEN), ("5", "RBAC角色", AMBER)]):
        x = 0.82 + i * 1.50
        add_rect(slide, x, 5.10, 1.24, 0.78, fill=PANEL, line=GRID, lw=0.6)
        add_text(slide, value, x + 0.10, 5.22, 0.44, 0.30, size=20, color=col, bold=True, font=MONO)
        add_text(slide, label, x + 0.56, 5.28, 0.58, 0.28, size=8, color=MUTED)
    # Right-side roadmap visual
    add_line(slide, 8.60, 1.95, 8.60, 5.52, color=GRID, width=1.6)
    milestones = [("07.29", "智能执行", CYAN), ("08.10", "多端部署", PURPLE), ("08.16", "需求对齐", AMBER), ("08.30+", "功能补全", GREEN)]
    for i, (date, label, col) in enumerate(milestones):
        y = 2.06 + i * 1.12
        add_circle(slide, 8.42, y, 0.36, fill=col, line=col, lw=0.5)
        add_text(slide, date, 9.02, y + 0.01, 1.15, 0.22, size=11, color=col, bold=True, font=MONO)
        add_text(slide, label, 10.30, y + 0.01, 1.60, 0.22, size=12, color=TEXT, bold=True)
    draw_shield(slide, 10.18, 4.88, 1.24, 1.40, CYAN)
    add_text(slide, "BUILD\n→ DEMO", 10.31, 5.33, 0.98, 0.34, size=9, color=TEXT, bold=True, align=PP_ALIGN.CENTER, font=MONO)
    add_text(slide, "WORK REPORT  ·  REDTEAM-PLATFORM", 7.55, 6.82, 4.62, 0.18, size=8, color=MUTED, align=PP_ALIGN.RIGHT, font=MONO)

    # 2. Timeline
    slide = prs.slides.add_slide(blank); add_dark_bg(slide); add_header(slide, "01 / TIMELINE", "工作主线：从智能执行方案到功能闭环", 2)
    add_text(slide, "每次会议要求都转化为可验证的代码、页面或部署文档", 0.62, 1.46, 8.0, 0.28, size=13, color=MUTED)
    add_line(slide, 1.12, 2.10, 1.12, 6.30, color=GRID, width=2.0)
    items = [
        ("01", "7 月 29 日 · ReAct 能力设计", "明确 LLM-in-the-loop、证据历史、动态步骤与安全阀。", CYAN),
        ("02", "8 月 06 日 · 一键测试体验", "修正阶段状态，打通推荐 Playbook 与攻击执行跳转。", BLUE),
        ("03", "8 月 10–15 日 · C/S 与多端部署", "完成服务器/云笔电部署、JWT、RBAC、WebSocket 与 ARM64 脚本。", PURPLE),
        ("04", "8 月 16 日 · 会议要求对齐", "明确应用程序渗透测试、拓扑范围与 PPT + 演示汇报方式。", AMBER),
        ("05", "8 月 30 日–至今 · 功能补全", "围绕攻击编排、扫描分析、报告、捕获、权限和靶场持续完善。", GREEN),
    ]
    for i, (num, title, desc, col) in enumerate(items):
        y = 2.05 + i * 0.88
        add_circle(slide, 0.94, y + 0.01, 0.36, fill=col, line=col, lw=0.5)
        add_text(slide, num, 0.94, y + 0.12, 0.36, 0.14, size=8, color=BG, bold=True, align=PP_ALIGN.CENTER, font=MONO)
        add_text(slide, title, 1.62, y, 4.52, 0.23, size=12, color=TEXT, bold=True)
        add_text(slide, desc, 1.62, y + 0.30, 5.30, 0.32, size=10, color=MUTED)
    add_rect(slide, 8.08, 2.10, 4.35, 3.78, fill=PANEL, line=CYAN, lw=0.8)
    add_text(slide, "本阶段关注点", 8.38, 2.38, 1.80, 0.28, size=15, color=TEXT, bold=True)
    for i, (lab, col) in enumerate([("功能能跑", CYAN), ("链路串起来", PURPLE), ("结果看得懂", GREEN), ("部署可复现", AMBER)]):
        y = 3.00 + i * 0.62
        add_circle(slide, 8.42, y + 0.05, 0.18, fill=col, line=col, lw=0.3)
        add_text(slide, lab, 8.82, y, 2.70, 0.22, size=12, color=TEXT, bold=True)
    add_text(slide, "→ 为后续现场演示准备稳定的主线", 8.40, 5.45, 3.54, 0.24, size=10, color=CYAN, bold=True)

    # 3. Architecture / deployment
    slide = prs.slides.add_slide(blank); add_dark_bg(slide); add_header(slide, "02 / PLATFORM", "平台架构与部署：服务器集中能力，终端按角色使用", 3)
    add_text(slide, "完成从单机验证到服务器 + 多终端的 C/S 运行方式", 0.62, 1.46, 8.0, 0.28, size=13, color=MUTED)
    # server card
    add_rect(slide, 0.72, 2.08, 7.10, 3.72, fill=PANEL, line=CYAN, lw=0.9)
    add_text(slide, "服务器端", 1.02, 2.32, 1.20, 0.24, size=14, color=CYAN, bold=True)
    add_text(slide, "Node.js API + SQLite + LLM + Docker/Podman", 1.02, 2.66, 5.82, 0.24, size=12, color=TEXT, bold=True, font=MONO)
    layers = [("API", "REST / JWT", CYAN), ("DATA", "SQLite 单文件", GREEN), ("AI", "LLM 分析与生成", PURPLE), ("TOOLS", "8 个镜像 · 30 个可执行工具", AMBER)]
    for i, (lab, desc, col) in enumerate(layers):
        x = 1.02 + (i % 2) * 3.15
        y = 3.30 + (i // 2) * 1.02
        add_rect(slide, x, y, 2.72, 0.72, fill=PANEL_2, line=col, lw=0.7)
        add_text(slide, lab, x + 0.18, y + 0.13, 0.72, 0.22, size=11, color=col, bold=True, font=MONO)
        add_text(slide, desc, x + 0.96, y + 0.15, 1.54, 0.20, size=9, color=MUTED)
    # client card
    add_rect(slide, 8.28, 2.08, 4.38, 3.72, fill=PANEL, line=PURPLE, lw=0.9)
    add_text(slide, "终端侧", 8.58, 2.32, 1.20, 0.24, size=14, color=PURPLE, bold=True)
    add_text(slide, "Qt 5.15.3 / C++17", 8.58, 2.68, 2.72, 0.24, size=12, color=TEXT, bold=True, font=MONO)
    for i, (lab, col) in enumerate([("管理员", CYAN), ("教师", PURPLE), ("学生", GREEN), ("操作员", AMBER), ("观察者", BLUE)]):
        add_chip(slide, lab, 8.58 + (i % 2) * 1.58, 3.26 + (i // 2) * 0.55, 1.30, col, size=8)
    add_text(slide, "角色化页面 · 服务器地址可配置 · WebSocket 实时状态", 8.58, 4.92, 3.58, 0.46, size=10, color=MUTED)
    add_arrow(slide, 7.90, 3.74, 0.35, 0.20, color=CYAN, transparency=10)
    # bottom checks
    add_rect(slide, 0.72, 6.18, 11.94, 0.42, fill=BG, line=GRID, lw=0.6)
    add_text(slide, "已验证：JWT 登录与刷新  ·  RBAC 路由守卫  ·  WebSocket 端点  ·  服务器/云笔电启动脚本", 0.92, 6.28, 11.55, 0.18, size=10, color=CYAN, bold=True, align=PP_ALIGN.CENTER)

    # 4. Scan workflow
    slide = prs.slides.add_slide(blank); add_dark_bg(slide); add_header(slide, "03 / SCAN", "扫描工作流：从目标识别到攻击方案推荐", 4)
    add_text(slide, "扫描页面从“任务列表”升级为“发现 + 结果 + 推荐”的工作台", 0.62, 1.46, 9.0, 0.28, size=13, color=MUTED)
    # left features
    add_rect(slide, 0.72, 2.04, 4.18, 4.02, fill=PANEL, line=CYAN, lw=0.8)
    add_card_title(slide, "新增应用程序扫描能力", 1.02, 2.30, 3.20, CYAN)
    scan_items = [("whatweb + httpx", "应用技术栈发现", CYAN), ("arjun + ffuf", "API 端点模糊测试", PURPLE), ("nuclei 模板", "认证机制审计", GREEN)]
    for i, (tool, desc, col) in enumerate(scan_items):
        y = 2.92 + i * 0.72
        add_circle(slide, 1.02, y + 0.05, 0.22, fill=col, line=col, lw=0.3)
        add_text(slide, tool, 1.42, y, 1.70, 0.20, size=10, color=TEXT, bold=True, font=MONO)
        add_text(slide, desc, 1.42, y + 0.26, 2.28, 0.18, size=9, color=MUTED)
    add_line(slide, 1.02, 5.20, 4.52, 5.20, color=GRID, width=0.6)
    add_text(slide, "目标分类", 1.02, 5.43, 0.90, 0.20, size=10, color=MUTED, bold=True)
    for i, lab in enumerate(["Web 应用", "REST API", "GraphQL", "SPA"]):
        add_chip(slide, lab, 1.98 + i * 0.66, 5.39, 0.58, [CYAN, PURPLE, AMBER, GREEN][i], size=7)
    # right workflow
    add_rect(slide, 5.28, 2.04, 7.38, 4.02, fill=PANEL, line=PURPLE, lw=0.8)
    add_text(slide, "扫描结果 → 推荐 Playbook → 预览 → 执行", 5.58, 2.30, 5.92, 0.28, size=15, color=TEXT, bold=True)
    wf = [("发现", "端口 / 技术栈 / 路径", CYAN), ("分析", "风险与攻击面", BLUE), ("推荐", "匹配现有方案", PURPLE), ("预览", "查看步骤后执行", GREEN)]
    for i, (title, desc, col) in enumerate(wf):
        x = 5.58 + i * 1.68
        add_circle(slide, x, 3.26, 0.46, fill=col, line=col, lw=0.4)
        add_text(slide, str(i + 1), x, 3.41, 0.46, 0.15, size=9, color=BG, bold=True, align=PP_ALIGN.CENTER, font=MONO)
        add_text(slide, title, x - 0.10, 3.88, 0.70, 0.20, size=11, color=TEXT, bold=True, align=PP_ALIGN.CENTER)
        add_text(slide, desc, x - 0.42, 4.19, 1.28, 0.36, size=8, color=MUTED, align=PP_ALIGN.CENTER)
        if i < 3: add_arrow(slide, x + 0.62, 3.42, 0.50, 0.18, color=col, transparency=10)
    add_rect(slide, 5.58, 5.12, 6.48, 0.56, fill=PANEL_2, line=GRID, lw=0.6)
    add_text(slide, "交互优化：点击推荐项先查看，确认后通过“前往执行”跳转。", 5.82, 5.29, 6.00, 0.18, size=10, color=CYAN, bold=True, align=PP_ALIGN.CENTER)

    # 5. ReAct / attack execution
    slide = prs.slides.add_slide(blank); add_dark_bg(slide); add_header(slide, "04 / ATTACK", "攻击执行与 ReAct：让执行过程可解释、可调整", 5)
    add_text(slide, "每一步工具执行后，LLM 基于证据决定继续、调整、插入、转向或停止", 0.62, 1.46, 10.0, 0.28, size=13, color=MUTED)
    # loop visual
    add_rect(slide, 0.72, 2.06, 7.20, 4.03, fill=PANEL, line=RED, lw=0.8)
    add_text(slide, "ReAct 执行闭环", 1.02, 2.34, 1.80, 0.24, size=14, color=TEXT, bold=True)
    loop = [("执行工具", CYAN), ("记录证据", BLUE), ("LLM 分析", PURPLE), ("调整计划", RED)]
    coords = [(1.18, 3.23), (3.32, 3.23), (3.32, 4.58), (1.18, 4.58)]
    for i, ((lab, col), (x, y)) in enumerate(zip(loop, coords)):
        add_rect(slide, x, y, 1.62, 0.62, fill=PANEL_2, line=col, lw=0.8)
        add_text(slide, lab, x, y + 0.19, 1.62, 0.20, size=11, color=TEXT, bold=True, align=PP_ALIGN.CENTER)
    add_arrow(slide, 2.83, 3.42, 0.38, 0.18, color=CYAN, transparency=10)
    add_arrow(slide, 4.00, 4.00, 0.18, 0.40, color=BLUE, transparency=10)
    add_arrow(slide, 2.83, 4.80, 0.38, 0.18, color=RED, transparency=10)
    add_arrow(slide, 1.72, 4.00, 0.18, 0.40, color=PURPLE, transparency=10)
    add_text(slide, "探索模式 + 启发式规则 + 失败恢复", 1.02, 5.58, 5.60, 0.22, size=10, color=RED, bold=True, font=MONO)
    # right bullets
    add_rect(slide, 8.28, 2.06, 4.38, 4.03, fill=PANEL, line=PURPLE, lw=0.8)
    add_text(slide, "本阶段完成的执行增强", 8.58, 2.34, 2.80, 0.24, size=14, color=TEXT, bold=True)
    bullets = ["支持 adjust / insert / pivot / stop 等动作", "HTTP、SQL 错误、登录页触发下一步建议", "失败工具黑名单与调用上限，避免死循环", "执行详情左右分栏：AI 推理 + 步骤 + 攻击证据", "三类攻击标签：窃取 / 篡改欺骗 / 关键设备夺控"]
    for i, t in enumerate(bullets):
        add_circle(slide, 8.62, 2.94 + i * 0.54, 0.14, fill=[CYAN, PURPLE, RED, GREEN, AMBER][i], line=[CYAN, PURPLE, RED, GREEN, AMBER][i], lw=0.3)
        add_text(slide, t, 8.92, 2.86 + i * 0.54, 3.38, 0.32, size=10, color=TEXT)

    # 6. Campaign orchestration
    slide = prs.slides.add_slide(blank); add_dark_bg(slide); add_header(slide, "05 / CAMPAIGN", "攻击战役编排：从单个 Playbook 到阶段化执行", 6)
    add_text(slide, "新增 Campaign 模型，支持阶段、预案、产物和执行状态的统一管理", 0.62, 1.46, 9.0, 0.28, size=13, color=MUTED)
    # phase chain
    phases = [("01", "数据抵近窃取", "凭据 / 文件 / 数据", CYAN), ("02", "信息篡改欺骗", "DNS / 会话 / 误导", PURPLE), ("03", "关键设备夺控", "横向 / 控制 / 持久", RED)]
    for i, (num, title, desc, col) in enumerate(phases):
        x = 0.84 + i * 3.04
        add_rect(slide, x, 2.15, 2.54, 1.34, fill=PANEL, line=col, lw=0.9)
        add_circle(slide, x + 0.20, 2.40, 0.36, fill=col, line=col, lw=0.4)
        add_text(slide, num, x + 0.20, 2.52, 0.36, 0.14, size=8, color=BG, bold=True, align=PP_ALIGN.CENTER, font=MONO)
        add_text(slide, title, x + 0.70, 2.38, 1.50, 0.24, size=12, color=TEXT, bold=True)
        add_text(slide, desc, x + 0.20, 2.98, 2.12, 0.20, size=9, color=MUTED)
        if i < 2: add_arrow(slide, x + 2.60, 2.73, 0.38, 0.18, color=col, transparency=10)
    # lower capabilities
    add_rect(slide, 0.84, 4.10, 11.82, 1.90, fill=PANEL, line=GREEN, lw=0.8)
    columns = [("阶段内编排", "串行 / 并行 / 条件执行", CYAN), ("跨阶段产物", "凭据、IP、Token 等传递", PURPLE), ("运行控制", "启动 / 暂停 / 终止 / 报告", GREEN), ("前端承载", "战役列表、阶段指示、产物面板", AMBER)]
    for i, (title, desc, col) in enumerate(columns):
        x = 1.14 + i * 2.86
        add_circle(slide, x, 4.52, 0.28, fill=col, line=col, lw=0.4)
        add_text(slide, title, x + 0.46, 4.46, 1.72, 0.20, size=11, color=TEXT, bold=True)
        add_text(slide, desc, x + 0.02, 4.96, 2.32, 0.42, size=9, color=MUTED)
    add_rect(slide, 0.84, 6.30, 11.82, 0.34, fill=BG, line=GRID, lw=0.6)
    add_text(slide, "阶段执行引擎与 Campaign 页面已搭建，后续继续做端到端联调和演示打磨。", 1.08, 6.38, 11.30, 0.18, size=10, color=CYAN, bold=True, align=PP_ALIGN.CENTER)

    # 7. Evaluation, capture and report
    slide = prs.slides.add_slide(blank); add_dark_bg(slide); add_header(slide, "06 / EVALUATION", "测试评估与报告：从执行证据到可交付结果", 7)
    add_text(slide, "补齐报告预览、导出和网络数据捕获入口，形成评估结果输出链路", 0.62, 1.46, 9.0, 0.28, size=13, color=MUTED)
    nodes = [("执行记录", "run / evidence", CYAN), ("捕获分析", "tcpdump / tshark", PURPLE), ("报告生成", "结构化模板", AMBER), ("多格式交付", "DOCX / PDF / HTML", GREEN)]
    for i, (title, desc, col) in enumerate(nodes):
        x = 0.88 + i * 3.02
        add_rect(slide, x, 2.20, 2.36, 1.05, fill=PANEL, line=col, lw=0.8)
        add_text(slide, title, x + 0.18, 2.43, 1.70, 0.22, size=12, color=TEXT, bold=True)
        add_text(slide, desc, x + 0.18, 2.78, 1.92, 0.18, size=9, color=MUTED, font=MONO)
        if i < 3: add_arrow(slide, x + 2.44, 2.62, 0.42, 0.18, color=col, transparency=10)
    # lower cards
    lower = [(0.88, "报告预览", "JSON 源码 ↔ HTML 预览\n支持模板化内容查看", CYAN), (4.02, "报告导出", "DOCX / PDF / HTML\n可直接在 WPS 中打开", GREEN), (7.16, "网络捕获", "任务创建、启动、停止\nPCAP 分析与协议统计", PURPLE), (10.30, "待联调", "实际抓包环境\n继续验证稳定性", AMBER)]
    for x, title, desc, col in lower:
        add_rect(slide, x, 4.12, 2.36, 1.56, fill=PANEL, line=col, lw=0.7)
        add_text(slide, title, x + 0.18, 4.38, 1.82, 0.22, size=12, color=TEXT, bold=True)
        add_text(slide, desc, x + 0.18, 4.82, 1.92, 0.54, size=10, color=MUTED)
    add_rect(slide, 0.88, 6.22, 11.78, 0.40, fill=BG, line=GRID, lw=0.6)
    add_text(slide, "汇报价值：不只展示“跑过工具”，还能展示证据、分析、评分与最终报告。", 1.12, 6.32, 11.30, 0.18, size=10, color=GREEN, bold=True, align=PP_ALIGN.CENTER)

    # 8. Assets: payload, knowledge, playbooks
    slide = prs.slides.add_slide(blank); add_dark_bg(slide); add_header(slide, "07 / ASSETS", "攻击资产沉淀：载荷库、知识图谱与通用 Playbook", 8)
    add_text(slide, "把一次攻击执行中的工具、载荷、知识和证据沉淀为可复用资产", 0.62, 1.46, 9.0, 0.28, size=13, color=MUTED)
    add_rect(slide, 0.72, 2.08, 3.75, 3.90, fill=PANEL, line=CYAN, lw=0.8)
    add_text(slide, "载荷库", 1.02, 2.34, 1.20, 0.24, size=14, color=TEXT, bold=True)
    for i, (lab, col) in enumerate([("分类 / 搜索", CYAN), ("详情 / 命令", BLUE), ("绑定 Playbook", PURPLE), ("AI 生成 / 审核", GREEN)]):
        y = 2.98 + i * 0.62
        add_circle(slide, 1.04, y + 0.06, 0.16, fill=col, line=col, lw=0.3)
        add_text(slide, lab, 1.38, y, 2.32, 0.22, size=11, color=TEXT, bold=True)
    add_chip(slide, "候选 → 审核 → 可用", 1.02, 5.34, 1.82, GREEN, size=8)
    add_rect(slide, 4.84, 2.08, 3.75, 3.90, fill=PANEL, line=PURPLE, lw=0.8)
    add_text(slide, "知识增强", 5.14, 2.34, 1.40, 0.24, size=14, color=TEXT, bold=True)
    add_text(slide, "OWASP Top 10 + MITRE 技术\n注入 ReAct 与 Playbook 生成上下文", 5.14, 2.82, 2.95, 0.62, size=12, color=MUTED)
    add_chip(slide, "技术栈识别", 5.14, 3.86, 1.08, CYAN, size=8)
    add_chip(slide, "攻击面分析", 6.36, 3.86, 1.08, PURPLE, size=8)
    add_chip(slide, "策略推荐", 5.14, 4.38, 1.08, AMBER, size=8)
    add_chip(slide, "证据约束", 6.36, 4.38, 1.08, GREEN, size=8)
    add_rect(slide, 8.96, 2.08, 3.70, 3.90, fill=PANEL, line=AMBER, lw=0.8)
    add_text(slide, "Playbook 通用化", 9.26, 2.34, 2.12, 0.24, size=14, color=TEXT, bold=True)
    add_text(slide, "从靶场路径依赖，转向\n基于扫描发现的应用目标适配", 9.26, 2.82, 2.90, 0.62, size=12, color=MUTED)
    for i, lab in enumerate(["Web 应用", "REST API", "GraphQL", "SPA"]):
        add_chip(slide, lab, 9.26 + (i % 2) * 1.28, 3.86 + (i // 2) * 0.52, 1.10, [CYAN, PURPLE, AMBER, GREEN][i], size=8)
    add_text(slide, "27 容器工具 + 3 虚拟工具 · 8 类镜像", 9.26, 5.34, 2.90, 0.18, size=9, color=AMBER, bold=True, font=MONO)

    # 9. Management and operations
    slide = prs.slides.add_slide(blank); add_dark_bg(slide); add_header(slide, "08 / OPERATIONS", "系统管理与运维：让平台可管理、可部署、可复现", 9)
    add_text(slide, "围绕管理员、教师、学生和操作员的实际使用，补齐管理与部署配套", 0.62, 1.46, 9.0, 0.28, size=13, color=MUTED)
    ops = [("权限矩阵", "模块 × 角色读写权限\n支持管理员可视化配置", CYAN), ("用户管理", "用户增删改查\nJWT 登录与刷新", PURPLE), ("靶场管理", "启动 / 停止 / 自检\n验证工具链是否正常", GREEN), ("部署运维", "服务器与云笔电脚本\n自动更新与 ARM64 构建", AMBER)]
    for i, (title, desc, col) in enumerate(ops):
        x = 0.72 + (i % 2) * 6.00
        y = 2.10 + (i // 2) * 1.64
        add_rect(slide, x, y, 5.54, 1.30, fill=PANEL, line=col, lw=0.8)
        add_circle(slide, x + 0.24, y + 0.32, 0.42, fill=col, line=col, lw=0.4)
        add_text(slide, str(i + 1), x + 0.24, y + 0.45, 0.42, 0.14, size=9, color=BG, bold=True, align=PP_ALIGN.CENTER, font=MONO)
        add_text(slide, title, x + 0.90, y + 0.26, 1.52, 0.22, size=12, color=TEXT, bold=True)
        add_text(slide, desc, x + 0.90, y + 0.62, 3.94, 0.36, size=10, color=MUTED)
        add_status(slide, "已补齐", x + 4.72, y + 0.48, col, width=0.62)
    add_rect(slide, 0.72, 5.60, 11.54, 0.78, fill=PANEL_2, line=GRID, lw=0.6)
    add_text(slide, "部署验证链", 1.02, 5.85, 1.12, 0.20, size=11, color=CYAN, bold=True, font=MONO)
    chain = [("安装", CYAN), ("启动", PURPLE), ("健康检查", GREEN), ("登录验收", AMBER), ("多端访问", BLUE)]
    for i, (lab, col) in enumerate(chain):
        x = 2.36 + i * 1.78
        add_circle(slide, x, 5.80, 0.26, fill=col, line=col, lw=0.3)
        add_text(slide, lab, x + 0.38, 5.82, 1.00, 0.16, size=9, color=TEXT, bold=True)
        if i < 4: add_arrow(slide, x + 1.28, 5.86, 0.34, 0.14, color=col, transparency=15)

    # 10. Summary and next steps
    slide = prs.slides.add_slide(blank); add_dark_bg(slide); add_header(slide, "09 / SUMMARY", "阶段成果与下一步：形成可演示、可继续演进的版本", 10)
    add_text(slide, "这段时间的核心产出，是把“功能点”逐步串成“可展示的工作链路”", 0.62, 1.46, 10.0, 0.28, size=13, color=MUTED)
    add_rect(slide, 0.72, 2.04, 5.72, 3.98, fill=PANEL, line=GREEN, lw=0.8)
    add_text(slide, "已完成 / 已落地", 1.02, 2.34, 1.80, 0.26, size=16, color=TEXT, bold=True)
    done = ["C/S 架构、JWT、RBAC、WebSocket 与多端部署", "扫描页面分栏、应用目标识别、推荐与预览执行", "ReAct 动态决策、失败恢复、攻击证据展示", "Campaign 阶段化编排、产物传递与运行控制", "报告预览/导出、抓包入口、载荷库与知识增强", "系统管理、靶场管理、启动/更新/验收脚本"]
    for i, t in enumerate(done):
        add_circle(slide, 1.06, 2.94 + i * 0.45, 0.14, fill=GREEN, line=GREEN, lw=0.3)
        add_text(slide, t, 1.38, 2.86 + i * 0.45, 4.62, 0.28, size=10, color=TEXT)
    add_rect(slide, 6.72, 2.04, 5.94, 3.98, fill=PANEL, line=AMBER, lw=0.8)
    add_text(slide, "下一步重点", 7.02, 2.34, 1.50, 0.26, size=16, color=TEXT, bold=True)
    nexts = [("01", "一站式 Pipeline 页面", "把扫描→分析→生成→执行做成单页面闭环", PURPLE), ("02", "抓包与评估联调", "验证 tcpdump/tshark 在目标环境的稳定运行", CYAN), ("03", "拓扑与应用测试演示", "完善设备/软件展示，并准备应用目标演示", GREEN), ("04", "许可证与交付打磨", "明确加密锁替代方案，固化安装与演示脚本", AMBER)]
    for i, (num, title, desc, col) in enumerate(nexts):
        y = 2.90 + i * 0.68
        add_text(slide, num, 7.04, y, 0.36, 0.18, size=10, color=col, bold=True, font=MONO)
        add_text(slide, title, 7.60, y - 0.02, 2.24, 0.20, size=11, color=TEXT, bold=True)
        add_text(slide, desc, 7.60, y + 0.25, 4.48, 0.18, size=9, color=MUTED)
    add_rect(slide, 0.72, 6.30, 11.94, 0.40, fill=BG, line=GRID, lw=0.6)
    add_text(slide, "汇报建议：先讲主线，再现场演示扫描 → 推荐 → 执行 → 报告。", 0.96, 6.40, 11.46, 0.18, size=11, color=CYAN, bold=True, align=PP_ALIGN.CENTER)

    return prs


if __name__ == "__main__":
    out_dir = Path(__file__).resolve().parents[1] / "docs" / "软件需规"
    out_dir.mkdir(parents=True, exist_ok=True)
    out_path = out_dir / "RedTeam-Platform_7.29以来工作汇报.pptx"
    make_progress_deck().save(str(out_path))
    print(out_path)

