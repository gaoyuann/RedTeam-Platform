from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pathlib import Path
import os


# ---- Theme ----
W, H = 13.333, 7.5
BG = RGBColor(247, 250, 253)
PANEL = RGBColor(255, 255, 255)
PANEL_2 = RGBColor(239, 246, 252)
TEXT = RGBColor(24, 45, 76)
MUTED = RGBColor(91, 108, 132)
CYAN = RGBColor(0, 146, 195)
BLUE = RGBColor(45, 108, 223)
PURPLE = RGBColor(103, 84, 196)
GREEN = RGBColor(28, 153, 111)
AMBER = RGBColor(200, 132, 15)
RED = RGBColor(210, 72, 91)
GRID = RGBColor(215, 226, 239)

FONT = "Microsoft YaHei"
MONO = "Consolas"


def rgb(hexstr):
    hexstr = hexstr.lstrip("#")
    return RGBColor(int(hexstr[0:2], 16), int(hexstr[2:4], 16), int(hexstr[4:6], 16))


def set_fill(shape, color, transparency=0):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.fill.transparency = transparency


def set_line(shape, color, width=1.0, transparency=0, dash=None):
    shape.line.color.rgb = color
    shape.line.width = Pt(width)
    shape.line.transparency = transparency
    if dash:
        shape.line.dash_style = dash


def add_rect(slide, x, y, w, h, fill=PANEL, radius=True, line=None, lw=1, transparency=0):
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shp = slide.shapes.add_shape(shape_type, Inches(x), Inches(y), Inches(w), Inches(h))
    set_fill(shp, fill, transparency)
    if line:
        set_line(shp, line, lw)
    else:
        set_line(shp, fill, 0.1, 100)
    return shp


def add_text(slide, text, x, y, w, h, size=14, color=TEXT, bold=False, align=PP_ALIGN.LEFT,
             font=FONT, valign=MSO_ANCHOR.MIDDLE, margin=0.04, italic=False):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(margin)
    tf.margin_right = Inches(margin)
    tf.margin_top = Inches(margin)
    tf.margin_bottom = Inches(margin)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    return box


def add_line(slide, x1, y1, x2, y2, color=GRID, width=1.0, dash=None, transparency=0):
    ln = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    set_line(ln, color, width, transparency, dash)
    return ln


def add_arrow(slide, x, y, w, h, color=CYAN, transparency=0):
    shp = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(x), Inches(y), Inches(w), Inches(h))
    set_fill(shp, color, transparency)
    set_line(shp, color, 0.1, 100)
    return shp


def add_circle(slide, x, y, d, fill=PANEL_2, line=CYAN, lw=1.2, transparency=0):
    shp = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d))
    set_fill(shp, fill, transparency)
    set_line(shp, line, lw)
    return shp


def add_chip(slide, label, x, y, w, color=CYAN, fill=None, size=9):
    fill = fill or color
    shp = add_rect(slide, x, y, w, 0.28, fill, radius=True, transparency=78)
    set_line(shp, color, 0.7, 10)
    add_text(slide, label, x, y + 0.01, w, 0.24, size=size, color=color, bold=True, align=PP_ALIGN.CENTER)
    return shp


def add_header(slide, section, title, page):
    add_text(slide, f"REDTEAM-PLATFORM  /  {section.upper()}", 0.55, 0.25, 5.4, 0.24, size=9, color=CYAN, bold=True, font=MONO)
    add_text(slide, title, 0.55, 0.58, 10.8, 0.54, size=24, color=TEXT, bold=True)
    add_line(slide, 0.55, 1.23, 12.78, 1.23, color=GRID, width=0.8)
    add_text(slide, f"{page:02d}  /  10", 11.75, 7.06, 1.0, 0.2, size=8, color=MUTED, align=PP_ALIGN.RIGHT, font=MONO)
    add_text(slide, "SECURE  ·  INTELLIGENT  ·  NATIVE", 0.58, 7.06, 3.6, 0.2, size=8, color=MUTED, font=MONO)


def add_dark_bg(slide, grid=True):
    bg = slide.background.fill
    bg.solid()
    bg.fore_color.rgb = BG
    if grid:
        # subtle grid
        for x in [0.45 + i * 0.55 for i in range(24)]:
            add_line(slide, x, 0.0, x, H, color=GRID, width=0.35, transparency=72)
        for y in [0.2 + i * 0.55 for i in range(14)]:
            add_line(slide, 0.0, y, W, y, color=GRID, width=0.35, transparency=72)
    # two soft accent blobs, kept light so diagrams stay readable on white
    c = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(10.65), Inches(-1.0), Inches(4.1), Inches(4.1))
    set_fill(c, PURPLE, 93); set_line(c, PURPLE, 0.1, 100)
    c2 = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(-1.4), Inches(5.55), Inches(3.4), Inches(3.4))
    set_fill(c2, CYAN, 94); set_line(c2, CYAN, 0.1, 100)


def add_metric(slide, value, label, x, y, w, color):
    add_rect(slide, x, y, w, 0.72, PANEL, radius=True, line=GRID, lw=0.6)
    add_text(slide, value, x + 0.14, y + 0.06, w * 0.43, 0.42, size=22, color=color, bold=True, font=MONO)
    add_text(slide, label, x + w * 0.43, y + 0.12, w * 0.5, 0.28, size=9, color=MUTED)


def add_bullet(slide, text, x, y, w, color=CYAN, size=12):
    add_circle(slide, x, y + 0.08, 0.12, fill=color, line=color, lw=0.3)
    add_text(slide, text, x + 0.2, y, w - 0.2, 0.30, size=size, color=TEXT)


def draw_shield(slide, x, y, w, h, color=CYAN):
    shp = slide.shapes.add_shape(MSO_SHAPE.HEXAGON, Inches(x), Inches(y), Inches(w), Inches(h))
    set_fill(shp, color, 82); set_line(shp, color, 1.5)
    inner = slide.shapes.add_shape(MSO_SHAPE.HEXAGON, Inches(x + 0.09), Inches(y + 0.10), Inches(w - 0.18), Inches(h - 0.20))
    set_fill(inner, BG, 10); set_line(inner, color, 1.0)
    add_line(slide, x + w * 0.32, y + h * 0.56, x + w * 0.48, y + h * 0.71, color=color, width=1.5)
    add_line(slide, x + w * 0.48, y + h * 0.71, x + w * 0.74, y + h * 0.39, color=color, width=1.5)


def make_presentation():
    prs = Presentation()
    prs.slide_width = Inches(W)
    prs.slide_height = Inches(H)
    blank = prs.slide_layouts[6]

    # 1 Cover
    slide = prs.slides.add_slide(blank); add_dark_bg(slide, grid=True)
    add_text(slide, "信息系统渗透智能化测试平台", 0.72, 1.18, 7.3, 0.80, size=29, bold=True)
    add_text(slide, "RedTeam-Platform", 0.75, 2.02, 5.6, 0.40, size=20, color=CYAN, bold=True, font=MONO)
    add_text(slide, "面向红队 / 渗透测试的智能化攻防演练平台\n集成扫描、攻击、评估全流程，支持 AI 驱动的 Playbook 自动生成与执行。", 0.75, 2.72, 6.25, 1.02, size=15, color=MUTED)
    add_chip(slide, "SCAN", 0.78, 4.16, 1.0, CYAN)
    add_chip(slide, "AI PLAYBOOK", 1.93, 4.16, 1.54, PURPLE)
    add_chip(slide, "CONTAINER", 3.64, 4.16, 1.40, GREEN)
    add_metric(slide, "45+", "预置 Playbook", 0.78, 5.42, 2.15, CYAN)
    add_metric(slide, "7", "功能模块", 3.10, 5.42, 1.75, PURPLE)
    add_metric(slide, "28", "子功能", 5.02, 5.42, 1.75, GREEN)
    # radar / attack visual
    cx, cy = 10.05, 3.36
    for d, tr in [(3.2, 90), (2.6, 91), (2.0, 92), (1.36, 94)]:
        c = add_circle(slide, cx - d/2, cy - d/2, d, fill=BG, line=CYAN, lw=0.8, transparency=tr)
    for a, b in [((8.54, 2.18), (11.56, 4.67)), ((8.73, 4.72), (11.30, 2.14)), ((8.10, 3.35), (11.96, 3.35))]:
        add_line(slide, *a, *b, color=PURPLE, width=0.8, transparency=20)
    for px, py, label, col in [(8.58, 2.12, "DISCOVER", CYAN), (11.52, 4.61, "EXPLOIT", RED), (11.26, 2.08, "ASSESS", GREEN), (8.10, 3.28, "TARGET", AMBER)]:
        add_circle(slide, px - 0.09, py - 0.09, 0.18, fill=col, line=col, lw=0.6)
        add_text(slide, label, px - 0.36, py + 0.15, 0.72, 0.18, size=7, color=col, bold=True, align=PP_ALIGN.CENTER, font=MONO)
    draw_shield(slide, 9.38, 2.60, 1.34, 1.52, CYAN)
    add_text(slide, "RED\nTEAM", 9.51, 3.03, 1.08, 0.42, size=10, color=TEXT, bold=True, align=PP_ALIGN.CENTER, font=MONO)
    add_text(slide, "PROJECT BRIEF  ·  PLATFORM OVERVIEW", 9.62, 6.85, 2.96, 0.18, size=8, color=MUTED, font=MONO, align=PP_ALIGN.RIGHT)

    # 2 Positioning
    slide = prs.slides.add_slide(blank); add_dark_bg(slide); add_header(slide, "01 / POSITIONING", "平台定位：把红队能力沉淀为可执行工作流", 2)
    add_text(slide, "面向红队 / 渗透测试的智能化攻防演练平台", 0.62, 1.48, 7.2, 0.30, size=13, color=MUTED)
    add_rect(slide, 0.72, 2.00, 5.38, 4.18, fill=PANEL, line=CYAN, lw=0.9)
    add_text(slide, "我们解决什么问题？", 1.02, 2.30, 2.60, 0.30, size=16, color=TEXT, bold=True)
    add_text(slide, "将目标发现、攻击验证、影响评估\n统一到一个可观测、可复用的演练闭环中。", 1.02, 2.78, 4.40, 0.72, size=14, color=MUTED)
    add_text(slide, "ONE PLATFORM  ·  ONE LOOP", 1.02, 3.72, 3.30, 0.20, size=9, color=CYAN, bold=True, font=MONO)
    flow = [("SCAN", CYAN), ("PLAYBOOK", PURPLE), ("ATTACK", RED), ("REPORT", GREEN)]
    for i, (lab, col) in enumerate(flow):
        x = 1.02 + i * 1.17
        add_chip(slide, lab, x, 4.18, 0.98 if lab != "PLAYBOOK" else 1.18, col, size=8)
        if i < 3: add_arrow(slide, x + (1.05 if lab != "PLAYBOOK" else 1.25), 4.24, 0.23, 0.14, color=col, transparency=15)
    add_text(slide, "结果：一次演练的发现与证据，可直接复用为下一次剧本。", 1.02, 5.20, 4.56, 0.42, size=11, color=TEXT)
    cards = [(6.55, 2.00, "统一入口", "任务、目标、结果与报告集中管理", CYAN),
             (9.15, 2.00, "智能编排", "LLM 基于扫描结果推荐攻击路径", PURPLE),
             (6.55, 4.18, "安全执行", "Docker / Podman 隔离运行工具", GREEN),
             (9.15, 4.18, "结果沉淀", "Playbook、知识图谱与报告资产化", AMBER)]
    for x, y, t, d, col in cards:
        add_rect(slide, x, y, 2.35, 1.66, fill=PANEL, line=col, lw=0.7)
        add_circle(slide, x + 0.18, y + 0.22, 0.34, fill=col, line=col, lw=0.5)
        add_text(slide, t, x + 0.66, y + 0.18, 1.52, 0.28, size=12, color=TEXT, bold=True)
        add_text(slide, d, x + 0.18, y + 0.72, 1.93, 0.66, size=10, color=MUTED)
    add_rect(slide, 6.55, 6.16, 4.95, 0.42, fill=BG, line=GRID, lw=0.6)
    add_text(slide, "可部署 · 可复用 · 可审计", 6.72, 6.24, 4.62, 0.18, size=11, color=CYAN, bold=True, align=PP_ALIGN.CENTER)

    # 3 Architecture
    slide = prs.slides.add_slide(blank); add_dark_bg(slide); add_header(slide, "02 / ARCHITECTURE", "四层架构：本地运行，能力可组合", 3)
    add_text(slide, "客户端、服务、数据与工具执行解耦，形成可控的单机闭环", 0.62, 1.46, 7.6, 0.30, size=13, color=MUTED)
    layers = [
        ("01", "Qt 5.15.3 / C++17", "原生桌面前端", "CMake · AppImage", CYAN),
        ("02", "Node.js Runtime", "本地 API 与任务编排", "localhost · WebSocket", PURPLE),
        ("03", "SQLite", "单文件数据与知识资产", "随软件分发 · 易备份", GREEN),
        ("04", "Podman / Docker", "工具隔离与执行引擎", "镜像离线打包", AMBER),
    ]
    y0 = 1.98
    for i, (num, name, desc, tag, col) in enumerate(layers):
        y = y0 + i * 1.12
        add_rect(slide, 0.82, y, 7.25, 0.86, fill=PANEL if i % 2 == 0 else PANEL_2, line=col, lw=0.8)
        add_text(slide, num, 1.03, y + 0.17, 0.48, 0.30, size=16, color=col, bold=True, font=MONO, align=PP_ALIGN.CENTER)
        add_text(slide, name, 1.70, y + 0.12, 2.80, 0.28, size=14, color=TEXT, bold=True, font=MONO)
        add_text(slide, desc, 1.70, y + 0.46, 2.95, 0.22, size=10, color=MUTED)
        add_chip(slide, tag, 5.63, y + 0.28, 1.96, col, size=8)
        if i < 3:
            add_arrow(slide, 3.82, y + 0.91, 0.56, 0.18, color=col, transparency=25)
    # side contract cards
    add_rect(slide, 8.52, 2.03, 3.72, 1.03, fill=PANEL, line=CYAN, lw=0.7)
    add_text(slide, "LOCAL-FIRST", 8.78, 2.25, 1.50, 0.22, size=11, color=CYAN, bold=True, font=MONO)
    add_text(slide, "后端绑定 localhost，数据不出机", 8.78, 2.59, 2.90, 0.23, size=11, color=TEXT)
    add_rect(slide, 8.52, 3.34, 3.72, 1.03, fill=PANEL, line=PURPLE, lw=0.7)
    add_text(slide, "EVENT-DRIVEN", 8.78, 3.56, 1.90, 0.22, size=11, color=PURPLE, bold=True, font=MONO)
    add_text(slide, "WebSocket 推送扫描与执行状态", 8.78, 3.90, 3.08, 0.23, size=11, color=TEXT)
    add_rect(slide, 8.52, 4.65, 3.72, 1.03, fill=PANEL, line=GREEN, lw=0.7)
    add_text(slide, "RBAC", 8.78, 4.87, 1.20, 0.22, size=11, color=GREEN, bold=True, font=MONO)
    add_text(slide, "admin / teacher / student 多角色", 8.78, 5.21, 3.05, 0.23, size=11, color=TEXT)
    add_text(slide, "数据流", 9.08, 6.20, 0.60, 0.2, size=9, color=MUTED, bold=True, font=MONO)
    add_line(slide, 9.82, 6.30, 11.84, 6.30, color=CYAN, width=1.1)
    for px, lab, col in [(9.80, "UI", CYAN), (10.50, "API", PURPLE), (11.18, "DB", GREEN), (11.84, "TOOL", AMBER)]:
        add_circle(slide, px - 0.07, 6.23, 0.14, fill=col, line=col, lw=0.4)
        add_text(slide, lab, px - 0.22, 6.47, 0.44, 0.16, size=7, color=col, align=PP_ALIGN.CENTER, font=MONO)

    # 4 Workflow
    slide = prs.slides.add_slide(blank); add_dark_bg(slide); add_header(slide, "03 / WORKFLOW", "核心工作流：从发现到评估的智能闭环", 4)
    add_text(slide, "每一步都有结果沉淀，下一步由数据驱动", 0.62, 1.46, 7.0, 0.30, size=13, color=MUTED)
    steps = [("01", "扫描目标", "资产 / 端口 / 漏洞", CYAN), ("02", "展示结果", "结构化发现", BLUE),
             ("03", "AI 生成", "匹配 Playbook", PURPLE), ("04", "攻击测试", "脚本化执行", RED), ("05", "评估报告", "评分 / 证据", GREEN)]
    xs = [0.72, 3.15, 5.58, 8.01, 10.44]
    for i, (num, title, desc, col) in enumerate(steps):
        add_rect(slide, xs[i], 2.32, 1.92, 1.54, fill=PANEL, line=col, lw=0.9)
        add_circle(slide, xs[i] + 0.16, 2.52, 0.34, fill=col, line=col, lw=0.5)
        add_text(slide, num, xs[i] + 0.16, 2.57, 0.34, 0.20, size=9, color=BG, bold=True, align=PP_ALIGN.CENTER, font=MONO)
        add_text(slide, title, xs[i] + 0.58, 2.50, 1.15, 0.30, size=13, color=TEXT, bold=True)
        add_text(slide, desc, xs[i] + 0.18, 3.08, 1.50, 0.26, size=10, color=MUTED)
        if i < 4: add_arrow(slide, xs[i] + 1.98, 2.96, 0.42, 0.22, color=col, transparency=10)
    # examples feed into AI
    add_text(slide, "典型数据流", 0.75, 4.54, 1.25, 0.24, size=11, color=CYAN, bold=True, font=MONO)
    add_rect(slide, 0.72, 4.95, 2.50, 1.08, fill=PANEL_2, line=GRID, lw=0.7)
    add_text(slide, "发现结果", 0.95, 5.14, 0.85, 0.22, size=10, color=MUTED, bold=True)
    add_text(slide, "443/tcp   ·   SQLi   ·   admin", 0.95, 5.47, 1.95, 0.20, size=10, color=TEXT, font=MONO)
    add_arrow(slide, 3.42, 5.28, 0.62, 0.24, color=PURPLE, transparency=10)
    add_rect(slide, 4.24, 4.95, 2.54, 1.08, fill=PANEL_2, line=PURPLE, lw=0.8)
    add_text(slide, "AI 决策", 4.50, 5.14, 0.90, 0.22, size=10, color=PURPLE, bold=True)
    add_text(slide, "目标 + 证据 + 约束", 4.50, 5.47, 1.95, 0.20, size=10, color=TEXT, font=MONO)
    add_arrow(slide, 7.00, 5.28, 0.62, 0.24, color=RED, transparency=10)
    add_rect(slide, 7.82, 4.95, 2.54, 1.08, fill=PANEL_2, line=RED, lw=0.8)
    add_text(slide, "Playbook", 8.08, 5.14, 0.90, 0.22, size=10, color=RED, bold=True)
    add_text(slide, "recon → exploit → verify", 8.08, 5.47, 2.05, 0.20, size=10, color=TEXT, font=MONO)
    add_arrow(slide, 10.42, 5.28, 0.62, 0.24, color=GREEN, transparency=10)
    add_rect(slide, 11.24, 4.95, 1.34, 1.08, fill=PANEL_2, line=GREEN, lw=0.8)
    add_text(slide, "REPORT", 11.42, 5.20, 0.96, 0.20, size=10, color=GREEN, bold=True, font=MONO, align=PP_ALIGN.CENTER)
    add_text(slide, "证据\n评分", 11.42, 5.48, 0.96, 0.38, size=10, color=TEXT, align=PP_ALIGN.CENTER)

    # 5 Modules
    slide = prs.slides.add_slide(blank); add_dark_bg(slide); add_header(slide, "04 / CAPABILITIES", "7 大功能模块，覆盖 28 个子功能", 5)
    add_text(slide, "从资源部署到知识沉淀，统一纳入任务、权限与审计体系", 0.62, 1.46, 8.0, 0.30, size=13, color=MUTED)
    modules = [
        ("01", "资源部署配置", "HEXSTRIKE · 沙箱 · LLM · 靶场 · 工具", CYAN),
        ("02", "漏洞利用想定", "Playbook · 脚本 · 执行引擎", PURPLE),
        ("03", "网络拓扑探测", "探测 · 编辑 · 展示 · 信息标注", BLUE),
        ("04", "脆弱性扫描", "nuclei · sqlmap · nmap · hydra", AMBER),
        ("05", "漏洞攻击测试", "窃取 · 篡改 · 夺控", RED),
        ("06", "测试评估", "流量捕获 · 数据分析评分", GREEN),
        ("07", "系统管理", "任务 · 用户 · 权限 · 报告 · 知识图谱", CYAN),
    ]
    positions = [(0.72, 2.05), (3.92, 2.05), (7.12, 2.05), (10.32, 2.05), (2.32, 4.33), (5.52, 4.33), (8.72, 4.33)]
    for (num, title, desc, col), (x, y) in zip(modules, positions):
        add_rect(slide, x, y, 2.55, 1.64, fill=PANEL, line=col, lw=0.85)
        add_circle(slide, x + 0.18, y + 0.18, 0.42, fill=col, line=col, lw=0.5)
        add_text(slide, num, x + 0.18, y + 0.28, 0.42, 0.15, size=9, color=BG, bold=True, align=PP_ALIGN.CENTER, font=MONO)
        add_text(slide, title, x + 0.74, y + 0.18, 1.52, 0.30, size=13, color=TEXT, bold=True)
        add_text(slide, desc, x + 0.20, y + 0.78, 2.15, 0.46, size=9, color=MUTED)
    add_rect(slide, 0.72, 6.42, 11.94, 0.38, fill=BG, line=GRID, lw=0.6)
    add_text(slide, "28 个子功能  ·  统一工作台  ·  角色化协同  ·  全程留痕", 0.92, 6.50, 11.55, 0.18, size=10, color=CYAN, bold=True, align=PP_ALIGN.CENTER)

    # 6 Scanning & attacks
    slide = prs.slides.add_slide(blank); add_dark_bg(slide); add_header(slide, "05 / TESTING", "脆弱性扫描 × 攻击测试：从发现到验证", 6)
    add_text(slide, "四类扫描建立攻击面，三类测试验证业务影响", 0.62, 1.46, 7.0, 0.30, size=13, color=MUTED)
    add_rect(slide, 0.68, 1.98, 5.72, 4.63, fill=PANEL, line=CYAN, lw=0.8)
    add_text(slide, "SCAN ENGINES", 0.98, 2.20, 2.0, 0.24, size=11, color=CYAN, bold=True, font=MONO)
    scans = [("NUCLEI", "系统漏洞模板扫描", "CVE / misconfig", CYAN), ("SQLMAP", "Web 注入验证", "SQLi / 数据库", PURPLE),
             ("NMAP", "端口与服务识别", "asset / service", BLUE), ("HYDRA", "弱口令检测", "credential", AMBER)]
    for i, (name, desc, tag, col) in enumerate(scans):
        y = 2.76 + i * 0.80
        add_circle(slide, 1.00, y, 0.38, fill=col, line=col, lw=0.4)
        add_text(slide, name[:2], 1.00, y + 0.10, 0.38, 0.14, size=9, color=BG, bold=True, align=PP_ALIGN.CENTER, font=MONO)
        add_text(slide, name, 1.58, y + 0.01, 1.20, 0.22, size=12, color=TEXT, bold=True, font=MONO)
        add_text(slide, desc, 2.86, y + 0.01, 1.65, 0.22, size=10, color=MUTED)
        add_chip(slide, tag, 4.78, y + 0.02, 1.18, col, size=8)
    add_line(slide, 6.66, 2.05, 6.66, 6.55, color=GRID, width=1.0)
    add_rect(slide, 6.95, 1.98, 5.72, 4.63, fill=PANEL, line=RED, lw=0.8)
    add_text(slide, "ATTACK TESTS", 7.24, 2.20, 2.0, 0.24, size=11, color=RED, bold=True, font=MONO)
    attacks = [("01", "数据抵近窃取", "验证敏感数据可达性", RED), ("02", "信息篡改欺骗", "验证完整性与业务影响", PURPLE), ("03", "关键设备夺控", "验证横向移动与控制面", AMBER)]
    for i, (num, title, desc, col) in enumerate(attacks):
        y = 2.76 + i * 1.06
        add_rect(slide, 7.24, y, 4.82, 0.80, fill=PANEL_2, line=col, lw=0.7)
        add_text(slide, num, 7.48, y + 0.20, 0.44, 0.24, size=14, color=col, bold=True, font=MONO)
        add_text(slide, title, 8.13, y + 0.11, 2.15, 0.24, size=12, color=TEXT, bold=True)
        add_text(slide, desc, 8.13, y + 0.44, 3.36, 0.18, size=9, color=MUTED)
    add_rect(slide, 7.24, 6.10, 4.82, 0.32, fill=BG, line=GRID, lw=0.5)
    add_text(slide, "发现 → 复现 → 证据 → 评分", 7.36, 6.16, 4.56, 0.16, size=9, color=GREEN, bold=True, font=MONO, align=PP_ALIGN.CENTER)

    # 7 AI Playbook
    slide = prs.slides.add_slide(blank); add_dark_bg(slide); add_header(slide, "06 / INTELLIGENCE", "AI 驱动的智能 Playbook：让发现直接变成行动", 7)
    add_text(slide, "基于扫描证据、知识图谱与安全边界，自动组合可审计的攻击剧本", 0.62, 1.46, 9.6, 0.30, size=13, color=MUTED)
    # left evidence
    add_rect(slide, 0.72, 2.12, 3.18, 3.74, fill=PANEL, line=CYAN, lw=0.8)
    add_text(slide, "SCAN FINDINGS", 0.98, 2.34, 1.80, 0.22, size=10, color=CYAN, bold=True, font=MONO)
    findings = [("asset", "10.10.4.21 / web"), ("port", "443 / https"), ("vuln", "CVE-2024-xxxx"), ("cred", "admin / weak pwd")]
    for i, (k, v) in enumerate(findings):
        y = 2.86 + i * 0.56
        add_text(slide, k.upper(), 0.98, y, 0.72, 0.20, size=9, color=MUTED, bold=True, font=MONO)
        add_text(slide, v, 1.82, y, 1.70, 0.20, size=10, color=TEXT, font=MONO)
        add_line(slide, 0.98, y + 0.30, 3.54, y + 0.30, color=GRID, width=0.5)
    add_chip(slide, "STRUCTURED EVIDENCE", 0.98, 5.32, 1.88, CYAN, size=8)
    # center LLM
    add_arrow(slide, 4.10, 3.62, 0.68, 0.28, color=PURPLE, transparency=12)
    add_circle(slide, 4.91, 2.60, 1.82, fill=PURPLE, line=PURPLE, lw=1.0, transparency=82)
    add_circle(slide, 5.10, 2.79, 1.44, fill=BG, line=PURPLE, lw=1.0)
    add_text(slide, "LLM", 5.10, 3.08, 1.44, 0.32, size=22, color=PURPLE, bold=True, align=PP_ALIGN.CENTER, font=MONO)
    add_text(slide, "匹配 · 规划 · 约束", 4.98, 4.60, 1.68, 0.24, size=10, color=MUTED, align=PP_ALIGN.CENTER)
    add_chip(slide, "KNOWLEDGE GRAPH", 4.72, 5.15, 2.18, PURPLE, size=8)
    # right playbook
    add_arrow(slide, 6.98, 3.62, 0.68, 0.28, color=GREEN, transparency=12)
    add_rect(slide, 7.82, 2.12, 4.76, 3.74, fill=PANEL, line=GREEN, lw=0.9)
    add_text(slide, "GENERATED PLAYBOOK", 8.10, 2.34, 2.40, 0.22, size=10, color=GREEN, bold=True, font=MONO)
    add_text(slide, "Web foothold → credential reuse →\ninternal service pivot", 8.10, 2.74, 3.95, 0.52, size=15, color=TEXT, bold=True)
    psteps = [("01", "recon_web_fingerprint", CYAN), ("02", "sqlmap_error_based", PURPLE), ("03", "netexec_scan", AMBER), ("04", "evidence_and_score", GREEN)]
    for i, (num, label, col) in enumerate(psteps):
        y = 3.60 + i * 0.43
        add_text(slide, num, 8.12, y, 0.34, 0.18, size=9, color=col, bold=True, font=MONO)
        add_text(slide, label, 8.62, y, 3.18, 0.18, size=10, color=MUTED, font=MONO)
    add_chip(slide, "可审计 · 可复用 · 可回滚", 8.10, 5.30, 2.10, GREEN, size=8)

    # 8 Native + container
    slide = prs.slides.add_slide(blank); add_dark_bg(slide); add_header(slide, "07 / DELIVERY", "兼容性与容器化：在受限环境中稳定交付", 8)
    add_text(slide, "Qt5 原生客户端 + 隔离工具链，兼顾兼容性、可迁移性与运维安全", 0.62, 1.46, 10.0, 0.30, size=13, color=MUTED)
    # left native
    add_rect(slide, 0.72, 2.05, 5.62, 3.58, fill=PANEL, line=CYAN, lw=0.9)
    add_text(slide, "客户端与系统兼容", 1.02, 2.32, 2.20, 0.30, size=16, color=TEXT, bold=True)
    add_text(slide, "Qt 5.15.3 + C++17", 1.02, 2.72, 2.70, 0.24, size=12, color=CYAN, bold=True, font=MONO)
    add_text(slide, "原生桌面体验，统一构建与分发", 1.02, 3.04, 3.40, 0.22, size=10, color=MUTED)
    # platform chips
    for i, (lab, col) in enumerate([("银河麒麟", CYAN), ("统信 UOS", BLUE), ("AppImage", PURPLE), ("deb / rpm", GREEN)]):
        add_chip(slide, lab, 1.02 + (i % 2) * 1.82, 3.68 + (i // 2) * 0.56, 1.54, col, size=8)
    add_line(slide, 4.60, 2.38, 5.80, 3.58, color=CYAN, width=0.8, transparency=25)
    add_line(slide, 4.60, 4.82, 5.80, 3.58, color=CYAN, width=0.8, transparency=25)
    add_circle(slide, 5.52, 3.28, 0.56, fill=CYAN, line=CYAN, lw=0.5)
    add_text(slide, "OS", 5.52, 3.46, 0.56, 0.14, size=9, color=BG, bold=True, align=PP_ALIGN.CENTER, font=MONO)
    # right containers
    add_rect(slide, 6.70, 2.05, 5.96, 3.58, fill=PANEL, line=GREEN, lw=0.9)
    add_text(slide, "容器化工具链", 7.00, 2.32, 1.80, 0.30, size=16, color=TEXT, bold=True)
    add_text(slide, "Docker / Podman", 7.00, 2.72, 2.70, 0.24, size=12, color=GREEN, bold=True, font=MONO)
    add_text(slide, "工具隔离运行，镜像可离线打包", 7.00, 3.04, 3.40, 0.22, size=10, color=MUTED)
    # container stack
    for i, (lab, col) in enumerate([("nuclei", CYAN), ("sqlmap", PURPLE), ("nmap", BLUE), ("hydra", AMBER)]):
        y = 3.64 + i * 0.42
        add_rect(slide, 7.00, y, 2.28, 0.28, fill=BG, line=col, lw=0.6)
        add_text(slide, lab, 7.18, y + 0.04, 1.20, 0.16, size=9, color=col, bold=True, font=MONO)
        add_text(slide, "isolated", 8.56, y + 0.04, 0.52, 0.16, size=8, color=MUTED, font=MONO, align=PP_ALIGN.RIGHT)
    add_text(slide, "镜像仓", 10.15, 4.00, 0.76, 0.18, size=9, color=MUTED, font=MONO, align=PP_ALIGN.CENTER)
    add_arrow(slide, 9.54, 4.04, 0.52, 0.20, color=GREEN, transparency=10)
    add_rect(slide, 10.72, 3.55, 1.36, 0.96, fill=BG, line=GREEN, lw=0.7)
    add_text(slide, "OFFLINE\nBUNDLE", 10.84, 3.82, 1.12, 0.42, size=10, color=GREEN, bold=True, align=PP_ALIGN.CENTER, font=MONO)
    # migration path
    add_text(slide, "离线迁移路径", 0.78, 6.08, 1.40, 0.22, size=10, color=CYAN, bold=True, font=MONO)
    for i, (lab, col) in enumerate([("构建镜像", CYAN), ("离线介质", PURPLE), ("目标环境", GREEN)]):
        x = 2.55 + i * 2.55
        add_circle(slide, x, 6.02, 0.40, fill=col, line=col, lw=0.4)
        add_text(slide, str(i + 1), x, 6.14, 0.40, 0.14, size=9, color=BG, bold=True, align=PP_ALIGN.CENTER, font=MONO)
        add_text(slide, lab, x + 0.53, 6.10, 1.14, 0.20, size=10, color=TEXT, bold=True)
        if i < 2: add_arrow(slide, x + 1.62, 6.12, 0.52, 0.18, color=col, transparency=15)

    # 9 Progress
    slide = prs.slides.add_slide(blank); add_dark_bg(slide); add_header(slide, "08 / ROADMAP", "当前进展与下一步计划", 9)
    add_text(slide, "先完成可运行骨架，再把智能能力做成可复用资产", 0.62, 1.46, 8.0, 0.30, size=13, color=MUTED)
    add_rect(slide, 0.72, 2.08, 5.70, 3.96, fill=PANEL, line=CYAN, lw=0.8)
    add_text(slide, "当前进展 · Phase 0", 1.02, 2.34, 2.10, 0.28, size=16, color=TEXT, bold=True)
    current = [("Qt5/C++ 项目骨架编译运行", CYAN), ("7 个一级模块导航占位", PURPLE), ("本地后端连接检测", GREEN), ("平台架构与功能清单确认", AMBER), ("扫描驱动工作流方案已定义", BLUE)]
    for i, (lab, col) in enumerate(current):
        y = 2.90 + i * 0.54
        add_circle(slide, 1.06, y + 0.06, 0.18, fill=col, line=col, lw=0.3)
        add_text(slide, lab, 1.44, y, 3.72, 0.22, size=11, color=TEXT)
        add_text(slide, "DONE", 5.45, y, 0.58, 0.20, size=8, color=col, bold=True, font=MONO, align=PP_ALIGN.RIGHT)
    add_rect(slide, 6.72, 2.08, 5.94, 3.96, fill=PANEL, line=PURPLE, lw=0.8)
    add_text(slide, "下一步计划", 7.02, 2.34, 1.50, 0.28, size=16, color=TEXT, bold=True)
    nexts = [("01", "数据库与核心数据迁移", "Phase 1：SQLite 数据模型与资产沉淀", PURPLE), ("02", "工具容器化", "Phase 2：工具镜像、执行引擎与离线包", GREEN), ("03", "Qt 主框架与工作流", "Phase 3/4：模块迁移与扫描驱动闭环", CYAN)]
    for i, (num, title, desc, col) in enumerate(nexts):
        y = 2.92 + i * 0.92
        add_rect(slide, 7.02, y, 5.30, 0.68, fill=PANEL_2, line=col, lw=0.7)
        add_text(slide, num, 7.26, y + 0.20, 0.42, 0.20, size=12, color=col, bold=True, font=MONO)
        add_text(slide, title, 7.86, y + 0.10, 1.78, 0.22, size=12, color=TEXT, bold=True)
        add_text(slide, desc, 7.86, y + 0.38, 3.98, 0.18, size=9, color=MUTED)
    add_rect(slide, 0.72, 6.34, 11.94, 0.40, fill=BG, line=GRID, lw=0.6)
    add_text(slide, "原则：能力组件化 · 过程可审计 · 结果可复用", 0.94, 6.44, 11.50, 0.18, size=10, color=CYAN, bold=True, align=PP_ALIGN.CENTER)

    # 10 Closing
    slide = prs.slides.add_slide(blank); add_dark_bg(slide, grid=True)
    add_text(slide, "谢谢", 0.80, 1.58, 4.0, 0.72, size=34, color=TEXT, bold=True)
    add_text(slide, "Q&A", 0.84, 2.40, 2.0, 0.38, size=22, color=CYAN, bold=True, font=MONO)
    add_text(slide, "信息系统渗透智能化测试平台", 0.84, 3.16, 4.6, 0.30, size=14, color=MUTED)
    add_line(slide, 0.84, 3.72, 4.78, 3.72, color=CYAN, width=1.4)
    for i, (lab, col) in enumerate([("DISCOVER", CYAN), ("AUTOMATE", PURPLE), ("ASSESS", GREEN)]):
        add_chip(slide, lab, 0.84 + i * 1.38, 4.08, 1.16, col, size=8)
    # closing radar visual
    cx, cy = 9.35, 3.62
    for d, col, tr in [(3.25, CYAN, 92), (2.45, PURPLE, 92), (1.64, GREEN, 91)]:
        add_circle(slide, cx - d/2, cy - d/2, d, fill=BG, line=col, lw=0.9, transparency=tr)
    add_line(slide, 7.70, 3.62, 11.00, 3.62, color=CYAN, width=0.8, transparency=20)
    add_line(slide, 9.35, 1.98, 9.35, 5.26, color=PURPLE, width=0.8, transparency=20)
    draw_shield(slide, 8.72, 2.86, 1.26, 1.44, GREEN)
    add_text(slide, "READY", 8.84, 3.35, 1.02, 0.18, size=10, color=TEXT, bold=True, align=PP_ALIGN.CENTER, font=MONO)
    add_text(slide, "REDTEAM-PLATFORM  ·  PROJECT BRIEF", 7.72, 6.66, 4.00, 0.22, size=9, color=MUTED, align=PP_ALIGN.CENTER, font=MONO)

    return prs


if __name__ == "__main__":
    # The repository's deliverables directory is writable in the shared workspace.
    out_dir = Path(__file__).resolve().parents[1] / "docs" / "软件需规"
    out_path = out_dir / "RedTeam-Platform_项目介绍.pptx"
    prs = make_presentation()
    prs.save(str(out_path))
    print(out_path)
